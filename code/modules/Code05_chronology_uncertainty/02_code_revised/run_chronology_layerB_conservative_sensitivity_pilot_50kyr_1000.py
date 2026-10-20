#!/usr/bin/env python3
"""Chronology Layer B conservative method-based sensitivity pilot.

Locked scope
------------
* Standard 50-kyr Global EV only; 1,000 realizations.
* Chronology perturbation only. Measurement values and record-wide z-scores
  remain fixed.
* Six independent chronology perturbation groups from the approved Layer B
  parameter table. All operational records in a group share one draw.
* ODP 1210 reuses the ODP 1209/Shatsky draw.
* ODP 690 receives an asymmetric post-K–Pg segment shift only; pre-K–Pg ages
  remain unchanged.
* Common K–Pg absolute-age, astronomical-target, and GPTS systematics are not
  included.
* Every realization re-bins measurements and dynamically reapplies n >= 2 and
  IQR > 0 eligibility; the observed record-bin mask is never reused.
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import math
import sys
import unicodedata
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


MASTER_SEED = 20260824
ITERATIONS = 1000
TOTAL_RETAINED_MEASUREMENTS = 10881
TOTAL_OPERATIONAL_RECORDS = 70
KPG_BOUNDARY_MA = 66.016

SCRIPT = Path(__file__).resolve()
WORK_ROOT = SCRIPT.parent.parent
DATA_DIR = WORK_ROOT / "01_data_revised"
FORMAL_DIR = DATA_DIR / "formal_datasets"
INVENTORY_DIR = WORK_ROOT / "03_results" / "EV_uncertainty_inventory"
PARAMETER_DIR = INVENTORY_DIR / "Chronology_LayerB_conservative_method_based_sensitivity"
PARAMETER_PATH = PARAMETER_DIR / "Chronology_LayerB_conservative_method_based_parameter_table.xlsx"
PRIMARY_FORMAL_DIR = INVENTORY_DIR / "Chronology_primary_sensitivity_formal_50kyr_10000"
PRIMARY_FORMAL_NPZ = PRIMARY_FORMAL_DIR / "Chronology_primary_formal_Global_EV_realizations_50kyr_10000.npz"
PRIMARY_PILOT_SCRIPT = WORK_ROOT / "02_code_revised" / "run_primary_chronology_sensitivity_pilot_50kyr_1000.py"
RESULT_DIR = INVENTORY_DIR / "Chronology_LayerB_conservative_sensitivity_pilot_50kyr_1000"
RESULT_DIR.mkdir(parents=True, exist_ok=True)

RAW_PATH = DATA_DIR / "retained_measurements_with_zscores_v4.csv"
IQR_PATH = DATA_DIR / "record_level_IQR_all_bin_widths_revised_v4_n2.csv"
S1_PATH = FORMAL_DIR / "Dataset_S1_revised_v6.xlsx"

GROUP_ORDER = [
    "CHR-P1-ODP1209-SHATSKY",
    "CHR-P2-U1403",
    "CHR-P3-SEYMOUR",
    "CHR-B1-ODP1262",
    "CHR-B2-ODP1267",
    "CHR-B3-ODP690-POSTKPG",
]
PRIMARY_SHIFT_KEYS = {
    "CHR-P1-ODP1209-SHATSKY": "ODP1209_shift_kyr",
    "CHR-P2-U1403": "U1403_shift_kyr",
    "CHR-P3-SEYMOUR": "Seymour_shift_kyr",
}
SEGMENT_GROUP = "CHR-B3-ODP690-POSTKPG"
CONTRAST_NAMES = [
    "Phase 1 - pre-escalation",
    "Phase 2 - pre-escalation",
    "Recovery - Phase 2",
]

PROTECTED_PATHS = [
    PARAMETER_PATH,
    PRIMARY_FORMAL_NPZ,
    RAW_PATH,
    IQR_PATH,
    S1_PATH,
    FORMAL_DIR / "Dataset_S2_revised.xlsx",
    FORMAL_DIR / "Dataset_S3_revised_v3.xlsx",
    FORMAL_DIR / "Dataset_S4_revised_v3.xlsx",
    FORMAL_DIR / "Dataset_S5_revised_v4.xlsx",
    FORMAL_DIR / "Dataset_S6_revised_v4.xlsx",
    WORK_ROOT / "05_manuscript_revised" / "Manuscript_working_stage3_sampling_density.docx",
    WORK_ROOT / "06_supplement_revised" / "Supplementary_materials_working_stage3_sampling_density.docx",
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def clean_id(value) -> str:
    return unicodedata.normalize("NFC", str(value)).strip()


def load_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot import {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def load_parameter_table(raw: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, list[str]]]:
    workbook = load_workbook(PARAMETER_PATH, read_only=True, data_only=True)

    sheet = workbook["Record parameters"]
    rows = list(sheet.iter_rows(min_row=4, values_only=True))
    headers = [clean_id(value) for value in rows[0]]
    records = pd.DataFrame(rows[1:], columns=headers).dropna(how="all")
    records["record_id"] = records.record_id.map(clean_id)
    records["perturbation_group_id"] = records.perturbation_group_id.fillna("").map(clean_id)
    records["active_flag"] = records.active_flag.astype(int)
    records["n_measurements"] = records.n_measurements.astype(int)

    sheet = workbook["Active groups"]
    rows = list(sheet.iter_rows(min_row=4, values_only=True))
    headers = [clean_id(value) for value in rows[0]]
    groups = pd.DataFrame(rows[1:], columns=headers).dropna(how="all")
    groups["perturbation_group_id"] = groups.perturbation_group_id.map(clean_id)
    for column in ["record_count", "measurement_count"]:
        groups[column] = groups[column].astype(int)
    for column in ["bound_min_kyr", "bound_max_kyr"]:
        groups[column] = groups[column].astype(float)

    if len(records) != 70 or records.record_id.nunique() != 70:
        raise RuntimeError("Layer B table must contain 70 unique records")
    if int(records.n_measurements.sum()) != TOTAL_RETAINED_MEASUREMENTS:
        raise RuntimeError("Layer B table does not reconcile to 10,881 measurements")
    active = records[records.active_flag.eq(1)].copy()
    if len(active) != 38 or int(active.n_measurements.sum()) != 8843:
        raise RuntimeError("Layer B active membership must contain 38 records / 8,843 measurements")
    if set(groups.perturbation_group_id) != set(GROUP_ORDER) or len(groups) != 6:
        raise RuntimeError("Layer B group inventory differs from the approved six-group design")

    raw_counts = raw.groupby("record_id").size()
    members: dict[str, list[str]] = {}
    for group_id in GROUP_ORDER:
        ids = active.loc[active.perturbation_group_id.eq(group_id), "record_id"].tolist()
        if not ids:
            raise RuntimeError(f"No active records mapped to {group_id}")
        missing = sorted(set(ids) - set(raw_counts.index))
        if missing:
            raise RuntimeError(f"{group_id} records missing from retained measurements: {missing}")
        group_row = groups.loc[groups.perturbation_group_id.eq(group_id)].iloc[0]
        actual_measurements = int(raw_counts.loc[ids].sum())
        if len(ids) != int(group_row.record_count) or actual_measurements != int(group_row.measurement_count):
            raise RuntimeError(f"{group_id} parameter/raw count mismatch")
        members[group_id] = ids

    if set(active.record_id) != set().union(*map(set, members.values())):
        raise RuntimeError("Active records do not map uniquely to the six perturbation groups")
    return records, groups, members


def draw_group_shifts(groups: pd.DataFrame, primary_npz) -> tuple[dict[str, np.ndarray], dict]:
    seed_sequence = np.random.SeedSequence(MASTER_SEED)
    children = seed_sequence.spawn(len(GROUP_ORDER))
    shifts: dict[str, np.ndarray] = {}
    seed_info = {}
    for group_id, child in zip(GROUP_ORDER, children):
        row = groups.loc[groups.perturbation_group_id.eq(group_id)].iloc[0]
        rng = np.random.default_rng(child)
        shifts[group_id] = rng.uniform(float(row.bound_min_kyr), float(row.bound_max_kyr), size=ITERATIONS)
        seed_info[group_id] = {
            "master_entropy": int(child.entropy),
            "spawn_key": list(child.spawn_key),
            "bit_generator": rng.bit_generator.__class__.__name__,
            "draw_min_kyr": float(shifts[group_id].min()),
            "draw_max_kyr": float(shifts[group_id].max()),
        }

    # The first three streams intentionally reproduce the primary formal run,
    # enabling paired Layer B-versus-primary comparisons.
    for group_id, npz_key in PRIMARY_SHIFT_KEYS.items():
        expected = np.asarray(primary_npz[npz_key][:ITERATIONS], dtype=float)
        if not np.array_equal(shifts[group_id], expected):
            raise RuntimeError(f"{group_id} shift draws do not reproduce the paired primary prefix")
    return shifts, seed_info


def critical_shift_keys(pilot, record_arrays, record_ids, shifts_ma: np.ndarray, segment_only: bool) -> tuple[np.ndarray, np.ndarray]:
    lower = pilot.BIN_SPEC["lower"]
    width = pilot.BIN_SPEC["width"]
    count = pilot.BIN_SPEC["count"]
    edges = lower + np.arange(count + 1) * width
    lo, hi = float(np.min(shifts_ma)), float(np.max(shifts_ma))
    critical = []
    for record_id in record_ids:
        ages = record_arrays[record_id][0]
        if segment_only:
            ages = ages[ages < KPG_BOUNDARY_MA]
        if len(ages) == 0:
            continue
        values = (edges[:, None] - ages[None, :]).ravel()
        critical.append(values[(values >= lo) & (values <= hi)])
    if not critical:
        return np.zeros(len(shifts_ma), dtype=int), np.array([], dtype=float)
    nonempty = [values for values in critical if len(values)]
    critical_values = np.unique(np.concatenate(nonempty)) if nonempty else np.array([], dtype=float)
    keys = np.searchsorted(critical_values, shifts_ma, side="right")
    return keys, critical_values


def segment_record_contributions(pilot, record_arrays, record_ids, shift_ma: float):
    n_bins = pilot.BIN_SPEC["count"]
    sums = np.zeros(n_bins, dtype=float)
    counts = np.zeros(n_bins, dtype=np.int16)
    diagnostics = {
        "record_bins_with_n_ge_2": 0,
        "eligible_nonzero_record_bins": 0,
        "zero_iqr_record_bins_excluded": 0,
        "invalid_iqr_record_bins": 0,
        "measurements_assigned_to_50kyr_domain": 0,
        "segment_shifted_measurements": 0,
        "age_inversions": 0,
        "segment_boundary_order_violations": 0,
    }
    for record_id in record_ids:
        ages, z_scores = record_arrays[record_id]
        post_mask = ages < KPG_BOUNDARY_MA
        perturbed = ages.copy()
        perturbed[post_mask] += shift_ma
        diagnostics["segment_shifted_measurements"] += int(post_mask.sum())

        original_order = np.argsort(ages, kind="mergesort")
        if np.any(np.diff(perturbed[original_order]) < -1e-12):
            diagnostics["age_inversions"] += 1
        if np.any(post_mask) and np.any(~post_mask):
            if float(np.max(perturbed[post_mask])) >= float(np.min(perturbed[~post_mask])) + 1e-12:
                diagnostics["segment_boundary_order_violations"] += 1

        assigned = pilot.assign_bins(perturbed)
        diagnostics["measurements_assigned_to_50kyr_domain"] += int(np.sum(assigned >= 0))
        for bin_index in np.unique(assigned[assigned >= 0]):
            values = z_scores[assigned == bin_index]
            if len(values) < 2:
                continue
            diagnostics["record_bins_with_n_ge_2"] += 1
            value = pilot.iqr_linear(values)
            if not math.isfinite(value) or value < 0:
                diagnostics["invalid_iqr_record_bins"] += 1
            elif value == 0:
                diagnostics["zero_iqr_record_bins_excluded"] += 1
            else:
                sums[int(bin_index)] += value
                counts[int(bin_index)] += 1
                diagnostics["eligible_nonzero_record_bins"] += 1
    return sums, counts, diagnostics


def cached_group_contributions(pilot, record_arrays, record_ids, shifts_kyr, group_id, segment_only=False):
    shifts_ma = np.asarray(shifts_kyr, dtype=float) / 1000.0
    keys, critical_values = critical_shift_keys(pilot, record_arrays, record_ids, shifts_ma, segment_only)
    unique_keys, inverse, frequencies = np.unique(keys, return_inverse=True, return_counts=True)
    unique_sums = np.empty((len(unique_keys), pilot.BIN_SPEC["count"]), dtype=float)
    unique_counts = np.empty((len(unique_keys), pilot.BIN_SPEC["count"]), dtype=np.int16)
    totals = Counter()
    for index, (key, frequency) in enumerate(zip(unique_keys, frequencies)):
        representative = int(np.flatnonzero(keys == key)[0])
        shift_ma = float(shifts_ma[representative])
        if segment_only:
            sums, counts, diagnostics = segment_record_contributions(pilot, record_arrays, record_ids, shift_ma)
        else:
            sums, counts, diagnostics = pilot.record_contributions(record_arrays, record_ids, shift_ma)
        unique_sums[index] = sums
        unique_counts[index] = counts
        totals.update({name: int(value) * int(frequency) for name, value in diagnostics.items()})
        if (index + 1) % 250 == 0 or index + 1 == len(unique_keys):
            print(f"{group_id}: {index + 1}/{len(unique_keys)} unique bin-assignment states", flush=True)
    return unique_sums[inverse], unique_counts[inverse], {
        "unique_assignment_states": int(len(unique_keys)),
        "critical_shift_boundaries": int(len(critical_values)),
        "diagnostic_totals": dict(totals),
    }


def bin_summary(ev: np.ndarray, ages: np.ndarray, observed: np.ndarray, label: str) -> pd.DataFrame:
    quantiles = np.nanquantile(ev, [0.025, 0.5, 0.975], axis=0)
    return pd.DataFrame({
        "scenario": label,
        "midpoint_Ma": ages,
        "observed_EV": observed,
        "simulation_median": quantiles[1],
        "simulation_2.5th": quantiles[0],
        "simulation_97.5th": quantiles[2],
        "envelope_width": quantiles[2] - quantiles[0],
        "observed_minus_simulation_median": observed - quantiles[1],
        "valid_realizations": np.sum(np.isfinite(ev), axis=0),
    })


def contrast_summary(contrasts: np.ndarray, label: str) -> pd.DataFrame:
    quantiles = np.nanquantile(contrasts, [0.025, 0.5, 0.975], axis=0)
    return pd.DataFrame([
        {
            "scenario": label,
            "contrast": name,
            "simulation_median": float(quantiles[1, index]),
            "simulation_2.5th": float(quantiles[0, index]),
            "simulation_97.5th": float(quantiles[2, index]),
            "interval_width": float(quantiles[2, index] - quantiles[0, index]),
            "valid_realizations": int(np.isfinite(contrasts[:, index]).sum()),
        }
        for index, name in enumerate(CONTRAST_NAMES)
    ])


def peak_midpoints(ev: np.ndarray, ages: np.ndarray) -> np.ndarray:
    peaks = np.full(len(ev), np.nan)
    for index in range(len(ev)):
        if np.any(np.isfinite(ev[index])):
            peaks[index] = ages[int(np.nanargmax(ev[index]))]
    return peaks


def peak_frequency(peaks: np.ndarray, label: str) -> pd.DataFrame:
    counts = pd.Series(peaks).value_counts(dropna=False).sort_index()
    return pd.DataFrame({
        "scenario": label,
        "peak_midpoint_Ma": counts.index,
        "frequency": counts.values.astype(int),
        "proportion": counts.values / len(peaks),
    })


def comparison_tables(layer_ev, layer_contrasts, layer_peaks, primary_npz, ages, observed):
    primary_1000 = np.asarray(primary_npz["primary_EV_50kyr"][:ITERATIONS], dtype=float)
    primary_10000 = np.asarray(primary_npz["primary_EV_50kyr"], dtype=float)
    primary_contrasts_1000 = np.asarray(primary_npz["primary_phase_contrasts"][:ITERATIONS], dtype=float)
    primary_contrasts_10000 = np.asarray(primary_npz["primary_phase_contrasts"], dtype=float)
    primary_peaks_1000 = np.asarray(primary_npz["primary_peak_midpoint_Ma"][:ITERATIONS], dtype=float)
    primary_peaks_10000 = np.asarray(primary_npz["primary_peak_midpoint_Ma"], dtype=float)

    if not np.array_equal(primary_npz["age_50kyr"], ages) or np.max(np.abs(primary_npz["observed_EV_50kyr"] - observed)) > 1e-12:
        raise RuntimeError("Primary formal result does not use the same observed baseline")

    scenario_arrays = {
        "Primary formal 10,000": primary_10000,
        "Primary matched 1,000": primary_1000,
        "Layer B pilot 1,000": layer_ev,
    }
    envelope_rows = []
    width_by_scenario = {}
    for label, values in scenario_arrays.items():
        q = np.quantile(values, [0.025, 0.975], axis=0)
        widths = q[1] - q[0]
        width_by_scenario[label] = widths
        envelope_rows.append({
            "scenario": label,
            "median_envelope_width": float(np.median(widths)),
            "maximum_envelope_width": float(np.max(widths)),
            "maximum_width_midpoint_Ma": float(ages[int(np.argmax(widths))]),
        })
    matched_change = width_by_scenario["Layer B pilot 1,000"] - width_by_scenario["Primary matched 1,000"]
    envelope_comparison = pd.DataFrame(envelope_rows)
    envelope_comparison["LayerB_minus_matched_primary_median_width"] = float(
        np.median(width_by_scenario["Layer B pilot 1,000"]) - np.median(width_by_scenario["Primary matched 1,000"])
    )
    envelope_comparison["LayerB_minus_matched_primary_maximum_width"] = float(
        np.max(width_by_scenario["Layer B pilot 1,000"]) - np.max(width_by_scenario["Primary matched 1,000"])
    )
    envelope_by_bin = pd.DataFrame({
        "midpoint_Ma": ages,
        "primary_matched_1000_width": width_by_scenario["Primary matched 1,000"],
        "LayerB_1000_width": width_by_scenario["Layer B pilot 1,000"],
        "LayerB_minus_primary_width": matched_change,
        "absolute_width_change": np.abs(matched_change),
    })

    formal_q = np.quantile(primary_contrasts_10000, [0.025, 0.5, 0.975], axis=0)
    matched_q = np.quantile(primary_contrasts_1000, [0.025, 0.5, 0.975], axis=0)
    layer_q = np.quantile(layer_contrasts, [0.025, 0.5, 0.975], axis=0)
    paired_difference = layer_contrasts - primary_contrasts_1000
    paired_q = np.quantile(paired_difference, [0.025, 0.5, 0.975], axis=0)
    contrast_rows = []
    for index, name in enumerate(CONTRAST_NAMES):
        contrast_rows.append({
            "contrast": name,
            "primary_formal_10000_median": float(formal_q[1, index]),
            "primary_formal_10000_2.5th": float(formal_q[0, index]),
            "primary_formal_10000_97.5th": float(formal_q[2, index]),
            "primary_matched_1000_median": float(matched_q[1, index]),
            "primary_matched_1000_2.5th": float(matched_q[0, index]),
            "primary_matched_1000_97.5th": float(matched_q[2, index]),
            "LayerB_1000_median": float(layer_q[1, index]),
            "LayerB_1000_2.5th": float(layer_q[0, index]),
            "LayerB_1000_97.5th": float(layer_q[2, index]),
            "paired_LayerB_minus_primary_median": float(paired_q[1, index]),
            "paired_difference_2.5th": float(paired_q[0, index]),
            "paired_difference_97.5th": float(paired_q[2, index]),
        })
    contrast_comparison = pd.DataFrame(contrast_rows)

    all_peak_ages = sorted(set(primary_peaks_1000[np.isfinite(primary_peaks_1000)]) | set(layer_peaks[np.isfinite(layer_peaks)]))
    primary_counts = Counter(primary_peaks_1000)
    layer_counts = Counter(layer_peaks)
    peak_comparison = pd.DataFrame([
        {
            "peak_midpoint_Ma": float(age),
            "primary_matched_1000_frequency": int(primary_counts[age]),
            "primary_matched_1000_proportion": float(primary_counts[age] / ITERATIONS),
            "LayerB_1000_frequency": int(layer_counts[age]),
            "LayerB_1000_proportion": float(layer_counts[age] / ITERATIONS),
            "LayerB_minus_primary_percentage_points": float(100 * (layer_counts[age] - primary_counts[age]) / ITERATIONS),
        }
        for age in all_peak_ages
    ])
    primary_mode = float(pd.Series(primary_peaks_1000).mode().iloc[0])
    layer_mode = float(pd.Series(layer_peaks).mode().iloc[0])
    tv_distance = 0.5 * float(sum(abs(layer_counts[age] - primary_counts[age]) / ITERATIONS for age in all_peak_ages))
    peak_metrics = {
        "primary_matched_1000_modal_peak_Ma": primary_mode,
        "LayerB_1000_modal_peak_Ma": layer_mode,
        "paired_realizations_with_changed_peak_fraction": float(np.mean(layer_peaks != primary_peaks_1000)),
        "total_variation_distance": tv_distance,
        "primary_formal_10000_modal_peak_Ma": float(pd.Series(primary_peaks_10000).mode().iloc[0]),
    }
    return envelope_comparison, envelope_by_bin, contrast_comparison, peak_comparison, peak_metrics


def main() -> None:
    pilot = load_module(PRIMARY_PILOT_SCRIPT, "primary_chronology_pilot_core")
    protected_before = {str(path): sha256(path) for path in PROTECTED_PATHS if path.exists()}

    raw = pd.read_csv(RAW_PATH).reset_index(drop=True)
    raw["record_id"] = raw.record_id.map(clean_id)
    if len(raw) != TOTAL_RETAINED_MEASUREMENTS or raw.record_id.nunique() != TOTAL_OPERATIONAL_RECORDS:
        raise RuntimeError("Final retained baseline must contain 10,881 measurements / 70 records")
    if raw.duplicated(["record_id", "age_Ma", "value", "sample_id"]).any():
        raise RuntimeError("Unexpected exact duplicate retained observations")

    parameter_records, parameter_groups, group_members = load_parameter_table(raw)
    record_arrays = pilot.prepare_record_arrays(raw)
    ages, observed, observed_counts, s1_diff, observed_zero_bins = pilot.validate_observed_baseline(raw, record_arrays)

    active_records = set(parameter_records.loc[parameter_records.active_flag.eq(1), "record_id"])
    fixed_records = [record_id for record_id in record_arrays if record_id not in active_records]
    if len(fixed_records) != 32:
        raise RuntimeError("Expected 32 fixed records")
    fixed_sums, fixed_counts, fixed_diagnostics = pilot.record_contributions(record_arrays, fixed_records, 0.0)

    baseline_group = {}
    for group_id in GROUP_ORDER:
        if group_id == SEGMENT_GROUP:
            baseline_group[group_id] = segment_record_contributions(pilot, record_arrays, group_members[group_id], 0.0)
        else:
            baseline_group[group_id] = pilot.record_contributions(record_arrays, group_members[group_id], 0.0)
    baseline_sums = fixed_sums + sum(baseline_group[group_id][0] for group_id in GROUP_ORDER)
    baseline_counts = fixed_counts + sum(baseline_group[group_id][1] for group_id in GROUP_ORDER)
    if not np.array_equal(baseline_counts, observed_counts):
        raise RuntimeError("Fixed + Layer B group baseline counts do not reproduce observed eligibility")
    if np.max(np.abs(baseline_sums / baseline_counts - observed)) > 1e-12:
        raise RuntimeError("Fixed + Layer B group baseline does not reproduce observed Global EV")

    primary_npz = np.load(PRIMARY_FORMAL_NPZ)
    shifts, seed_info = draw_group_shifts(parameter_groups, primary_npz)

    group_iteration_sums = {}
    group_iteration_counts = {}
    group_cache_diagnostics = {}
    for group_id in GROUP_ORDER:
        sums, counts, diagnostics = cached_group_contributions(
            pilot,
            record_arrays,
            group_members[group_id],
            shifts[group_id],
            group_id,
            segment_only=group_id == SEGMENT_GROUP,
        )
        group_iteration_sums[group_id] = sums
        group_iteration_counts[group_id] = counts
        group_cache_diagnostics[group_id] = diagnostics

    total_sums = np.broadcast_to(fixed_sums, (ITERATIONS, pilot.BIN_SPEC["count"])).copy()
    total_counts = np.broadcast_to(fixed_counts, (ITERATIONS, pilot.BIN_SPEC["count"])).copy()
    for group_id in GROUP_ORDER:
        total_sums += group_iteration_sums[group_id]
        total_counts += group_iteration_counts[group_id]
    ev = np.divide(total_sums, total_counts, out=np.full_like(total_sums, np.nan), where=total_counts > 0)
    contrasts = pilot.phase_contrasts(ev, ages)
    peaks = peak_midpoints(ev, ages)

    ev_summary = bin_summary(ev, ages, observed, "Chronology Layer B pilot 1,000")
    phase_summary = contrast_summary(contrasts, "Chronology Layer B pilot 1,000")
    peak_summary = peak_frequency(peaks, "Chronology Layer B pilot 1,000")

    eligibility = pd.DataFrame({
        "midpoint_Ma": ages,
        "observed_eligible_record_bins": observed_counts,
        "simulation_record_bins_min": np.min(total_counts, axis=0),
        "simulation_record_bins_median": np.median(total_counts, axis=0),
        "simulation_record_bins_max": np.max(total_counts, axis=0),
        "simulation_record_bins_2.5th": np.quantile(total_counts, 0.025, axis=0),
        "simulation_record_bins_97.5th": np.quantile(total_counts, 0.975, axis=0),
        "iterations_with_empty_bin": np.sum(total_counts == 0, axis=0),
    })
    total_eligible_record_bins = total_counts.sum(axis=1)
    eligibility_distribution = pd.DataFrame([{
        "metric": "Total eligible record-bins per realization",
        "minimum": int(total_eligible_record_bins.min()),
        "2.5th": float(np.quantile(total_eligible_record_bins, 0.025)),
        "median": float(np.median(total_eligible_record_bins)),
        "97.5th": float(np.quantile(total_eligible_record_bins, 0.975)),
        "maximum": int(total_eligible_record_bins.max()),
        "observed": int(observed_counts.sum()),
    }])

    envelope_comparison, envelope_by_bin, contrast_comparison, peak_comparison, peak_metrics = comparison_tables(
        ev, contrasts, peaks, primary_npz, ages, observed
    )

    widths = ev_summary.envelope_width.to_numpy(float)
    absolute_difference = np.abs(ev_summary.observed_EV - ev_summary.simulation_median).to_numpy(float)
    segment_diagnostics = group_cache_diagnostics[SEGMENT_GROUP]["diagnostic_totals"]
    invalid_iqrs = int(fixed_diagnostics["invalid_iqr_record_bins"])
    for diagnostics in group_cache_diagnostics.values():
        invalid_iqrs += int(diagnostics["diagnostic_totals"].get("invalid_iqr_record_bins", 0))
    anomaly = {
        "age_inversions": int(segment_diagnostics.get("age_inversions", 0)),
        "segment_boundary_order_violations": int(segment_diagnostics.get("segment_boundary_order_violations", 0)),
        "segment_discontinuity_interpretation": "ODP690 uses the approved intentional piecewise post-K–Pg shift; zero order violations means no unintended age inversion across the segment boundary.",
        "ODP690_post_KPg_measurements_shifted_per_realization": int(
            sum(np.sum(record_arrays[record_id][0] < KPG_BOUNDARY_MA) for record_id in group_members[SEGMENT_GROUP])
        ),
        "ODP690_shift_draw_min_kyr": float(shifts[SEGMENT_GROUP].min()),
        "ODP690_shift_draw_max_kyr": float(shifts[SEGMENT_GROUP].max()),
        "iterations_with_at_least_one_empty_bin": int(np.sum(np.any(total_counts == 0, axis=1))),
        "total_empty_bin_events": int(np.sum(total_counts == 0)),
        "invalid_or_missing_EV_trajectory_iterations": int(np.sum(~np.all(np.isfinite(ev), axis=1))),
        "nonfinite_EV_values": int(np.sum(~np.isfinite(ev))),
        "invalid_IQR_events": invalid_iqrs,
        "observed_total_eligible_record_bins": int(observed_counts.sum()),
        "simulated_total_eligible_record_bins_min": int(total_eligible_record_bins.min()),
        "simulated_total_eligible_record_bins_median": float(np.median(total_eligible_record_bins)),
        "simulated_total_eligible_record_bins_max": int(total_eligible_record_bins.max()),
        "maximum_absolute_per_bin_record_count_change": int(np.max(np.abs(total_counts - observed_counts[None, :]))),
        "observed_zero_IQR_record_bins_excluded": int(observed_zero_bins),
    }
    if anomaly["age_inversions"] or anomaly["segment_boundary_order_violations"] or anomaly["invalid_IQR_events"]:
        raise RuntimeError(f"Chronology Layer B pilot failed anomaly checks: {anomaly}")

    group_shift_table = pd.DataFrame({"iteration": np.arange(1, ITERATIONS + 1)})
    for group_id in GROUP_ORDER:
        group_shift_table[f"{group_id}_shift_kyr"] = shifts[group_id]

    group_parameters = parameter_groups.copy()
    group_parameters["actual_record_count"] = group_parameters.perturbation_group_id.map(
        lambda group_id: len(group_members[group_id])
    )
    raw_counts = raw.groupby("record_id").size()
    group_parameters["actual_measurement_count"] = group_parameters.perturbation_group_id.map(
        lambda group_id: int(raw_counts.loc[group_members[group_id]].sum())
    )
    group_parameters["random_seed"] = MASTER_SEED

    active_membership = parameter_records.loc[parameter_records.active_flag.eq(1), [
        "record_id", "category", "sheet", "n_measurements", "chronology_method_class", "actual_chronology_method",
        "scenario_role", "perturbation_group_id", "recommended_perturbation_structure", "bound_min_kyr",
        "bound_max_kyr", "distribution", "applicable_interval", "correlation_rule", "range_basis",
    ]].copy()

    realizations_wide = pd.DataFrame(ev, columns=[f"EV_{age:.2f}_Ma" for age in ages])
    realizations_wide.insert(0, "total_eligible_record_bins", total_eligible_record_bins)
    realizations_wide.insert(0, "peak_midpoint_Ma", peaks)
    realizations_wide.insert(0, "iteration", np.arange(1, ITERATIONS + 1))

    phase_realizations = pd.DataFrame(contrasts, columns=CONTRAST_NAMES)
    phase_realizations.insert(0, "iteration", np.arange(1, ITERATIONS + 1))

    primary_width = envelope_by_bin.primary_matched_1000_width.to_numpy(float)
    layer_width = envelope_by_bin.LayerB_1000_width.to_numpy(float)
    comparison_metrics = {
        "LayerB_minus_primary_matched_median_envelope_width": float(np.median(layer_width) - np.median(primary_width)),
        "LayerB_minus_primary_matched_maximum_envelope_width": float(np.max(layer_width) - np.max(primary_width)),
        "median_bin_specific_envelope_width_change": float(np.median(layer_width - primary_width)),
        "maximum_absolute_bin_specific_width_change": float(np.max(np.abs(layer_width - primary_width))),
        "maximum_absolute_width_change_midpoint_Ma": float(ages[int(np.argmax(np.abs(layer_width - primary_width)))]),
        **peak_metrics,
    }

    phase2_row = phase_summary.loc[phase_summary.contrast.eq("Phase 2 - pre-escalation")].iloc[0]
    metadata = {
        "analysis": "Chronology Layer B conservative method-based sensitivity pilot — standard 50-kyr Global EV",
        "iterations": ITERATIONS,
        "master_random_seed": MASTER_SEED,
        "random_streams": seed_info,
        "measurement_values_perturbed": False,
        "measurement_uncertainty_included": False,
        "common_absolute_timescale_systematic_included": False,
        "common_astronomical_target_systematic_included": False,
        "common_GPTS_systematic_included": False,
        "observed_record_bin_mask_reused": False,
        "bin_assignment": "49 standard 50-kyr half-open bins; 64.775 <= age < 67.225 Ma",
        "eligibility": "n_measurements >= 2 and linearly interpolated IQR > 0, recomputed every realization",
        "global_EV": "record-weighted arithmetic mean of dynamically eligible record-bin IQRs",
        "active_chronology_groups": 6,
        "active_operational_records": len(active_records),
        "active_measurements": int(parameter_records.loc[parameter_records.active_flag.eq(1), "n_measurements"].sum()),
        "active_measurement_coverage_fraction": 8843 / TOTAL_RETAINED_MEASUREMENTS,
        "active_measurement_coverage_percent": 100 * 8843 / TOTAL_RETAINED_MEASUREMENTS,
        "fixed_operational_records": len(fixed_records),
        "retained_operational_records": int(raw.record_id.nunique()),
        "retained_measurements": int(len(raw)),
        "Dataset_S1_observed_EV_max_absolute_validation_difference": s1_diff,
        "observed_EV_vs_simulation_median_max_absolute_difference": float(np.max(absolute_difference)),
        "maximum_difference_midpoint_Ma": float(ev_summary.iloc[int(np.argmax(absolute_difference))].midpoint_Ma),
        "envelope_median_width": float(np.median(widths)),
        "envelope_maximum_width": float(np.max(widths)),
        "maximum_width_midpoint_Ma": float(ev_summary.iloc[int(np.argmax(widths))].midpoint_Ma),
        "observed_peak_midpoint_Ma": float(ages[int(np.argmax(observed))]),
        "Phase2_minus_pre_median": float(phase2_row.simulation_median),
        "Phase2_minus_pre_2.5th": float(phase2_row["simulation_2.5th"]),
        "Phase2_minus_pre_97.5th": float(phase2_row["simulation_97.5th"]),
        "Phase2_stable_positive_interval": bool(phase2_row["simulation_2.5th"] > 0),
        "anomaly_diagnostics": anomaly,
        "primary_comparison": comparison_metrics,
    }

    # Save machine-readable outputs before workbook authoring.
    ev_summary.to_csv(RESULT_DIR / "Chronology_LayerB_pilot_50kyr_Global_EV_summary.csv", index=False)
    phase_summary.to_csv(RESULT_DIR / "Chronology_LayerB_pilot_phase_contrasts.csv", index=False)
    peak_summary.to_csv(RESULT_DIR / "Chronology_LayerB_pilot_peak_midpoint_frequency.csv", index=False)
    eligibility.to_csv(RESULT_DIR / "Chronology_LayerB_pilot_eligibility_by_bin.csv", index=False)
    eligibility_distribution.to_csv(RESULT_DIR / "Chronology_LayerB_pilot_eligibility_distribution.csv", index=False)
    group_shift_table.to_csv(RESULT_DIR / "Chronology_LayerB_pilot_group_shift_draws.csv", index=False)
    group_parameters.to_csv(RESULT_DIR / "Chronology_LayerB_pilot_group_parameters.csv", index=False)
    active_membership.to_csv(RESULT_DIR / "Chronology_LayerB_pilot_active_record_membership.csv", index=False)
    realizations_wide.to_csv(RESULT_DIR / "Chronology_LayerB_pilot_Global_EV_realizations_50kyr_1000.csv", index=False)
    phase_realizations.to_csv(RESULT_DIR / "Chronology_LayerB_pilot_phase_contrast_realizations_1000.csv", index=False)
    envelope_comparison.to_csv(RESULT_DIR / "Chronology_LayerB_vs_primary_envelope_comparison.csv", index=False)
    envelope_by_bin.to_csv(RESULT_DIR / "Chronology_LayerB_vs_primary_envelope_change_by_bin.csv", index=False)
    contrast_comparison.to_csv(RESULT_DIR / "Chronology_LayerB_vs_primary_phase_contrast_comparison.csv", index=False)
    peak_comparison.to_csv(RESULT_DIR / "Chronology_LayerB_vs_primary_peak_distribution_comparison.csv", index=False)
    (RESULT_DIR / "Chronology_LayerB_pilot_metadata.json").write_text(json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8")
    (RESULT_DIR / "Chronology_LayerB_pilot_group_cache_diagnostics.json").write_text(
        json.dumps(group_cache_diagnostics, indent=2, ensure_ascii=False), encoding="utf-8"
    )

    np.savez_compressed(
        RESULT_DIR / "Chronology_LayerB_pilot_complete_realizations_50kyr_1000.npz",
        age_50kyr=ages,
        observed_EV_50kyr=observed,
        LayerB_EV_50kyr=ev,
        LayerB_eligible_record_counts=total_counts,
        LayerB_phase_contrasts=contrasts,
        LayerB_peak_midpoint_Ma=peaks,
        total_eligible_record_bins=total_eligible_record_bins,
        **{f"shift_{group_id}": shifts[group_id] for group_id in GROUP_ORDER},
    )

    protected_after = {str(path): sha256(path) for path in PROTECTED_PATHS if path.exists()}
    protected_check = pd.DataFrame([
        {
            "path": path,
            "sha256_before": digest,
            "sha256_after": protected_after.get(path),
            "unchanged": digest == protected_after.get(path),
        }
        for path, digest in protected_before.items()
    ])
    protected_check.to_csv(RESULT_DIR / "protected_files_hash_check.csv", index=False)
    if not protected_check.unchanged.all():
        raise RuntimeError("At least one protected Dataset/document/source file changed during the pilot")

    manifest_paths = [RAW_PATH, IQR_PATH, S1_PATH, PARAMETER_PATH, PRIMARY_FORMAL_NPZ, SCRIPT]
    pd.DataFrame([
        {"file": path.name, "path": str(path), "sha256": sha256(path)} for path in manifest_paths
    ]).to_csv(RESULT_DIR / "input_source_hash_manifest.csv", index=False)

    print(json.dumps(metadata, indent=2, ensure_ascii=False))
    print("\nPHASE CONTRASTS")
    print(phase_summary.to_string(index=False))
    print("\nPRIMARY COMPARISON — ENVELOPE")
    print(envelope_comparison.to_string(index=False))
    print("\nPRIMARY COMPARISON — CONTRASTS")
    print(contrast_comparison.to_string(index=False))
    print("\nPRIMARY COMPARISON — PEAKS")
    print(peak_comparison.to_string(index=False))


if __name__ == "__main__":
    main()
