#!/usr/bin/env python3
"""Formal 10,000-iteration primary chronology sensitivity for 50-kyr Global EV.

The validated pilot's method, grouping, bounds, seed, binning, IQR definition,
and dynamic eligibility rules are unchanged.  The first 1,000 realizations are
explicitly compared with the validated pilot.  An expanded 1,000-realization
ODP 1210 diagnostic is kept separate from the primary result and uses the same
Shatsky shift as ODP 1209.
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import math
import shutil
import sys
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


MASTER_SEED = 20260824
ITERATIONS = 10000
EXPANDED_ITERATIONS = 1000
PREFIXES = (1000, 5000, 10000)

SCRIPT = Path(__file__).resolve()
WORK_ROOT = SCRIPT.parent.parent
DATA_DIR = WORK_ROOT / "01_data_revised"
FORMAL_DIR = DATA_DIR / "formal_datasets"
INVENTORY_DIR = WORK_ROOT / "03_results" / "EV_uncertainty_inventory"
PILOT_DIR = INVENTORY_DIR / "Chronology_primary_sensitivity_pilot_50kyr_1000"
RESULT_DIR = INVENTORY_DIR / "Chronology_primary_sensitivity_formal_50kyr_10000"
RESULT_DIR.mkdir(parents=True, exist_ok=True)

PILOT_SCRIPT = WORK_ROOT / "02_code_revised" / "run_primary_chronology_sensitivity_pilot_50kyr_1000.py"
PILOT_NPZ = PILOT_DIR / "Chronology_primary_Global_EV_realizations_50kyr_1000.npz"
RAW_PATH = DATA_DIR / "retained_measurements_with_zscores_v4.csv"
IQR_PATH = DATA_DIR / "record_level_IQR_all_bin_widths_revised_v4_n2.csv"
S1_PATH = FORMAL_DIR / "Dataset_S1_revised_v6.xlsx"
GROUPING_PATH = INVENTORY_DIR / "Chronology_shared_age_model_grouping_audit" / "Chronology_shared_age_model_grouping_audit.xlsx"

GROUP_ORDER = ["CHR-P1-ODP1209", "CHR-P2-U1403", "CHR-P3-SEYMOUR"]
CONTRAST_NAMES = ["Phase 1 - pre-escalation", "Phase 2 - pre-escalation", "Recovery - Phase 2"]
EXPANDED_GROUP_ID = "CHR-X1-ODP1210"

PROTECTED_PATHS = [
    GROUPING_PATH,
    S1_PATH,
    FORMAL_DIR / "Dataset_S2_revised.xlsx",
    FORMAL_DIR / "Dataset_S3_revised_v3.xlsx",
    FORMAL_DIR / "Dataset_S4_revised_v3.xlsx",
    FORMAL_DIR / "Dataset_S5_revised_v4.xlsx",
    FORMAL_DIR / "Dataset_S6_revised_v4.xlsx",
    WORK_ROOT / "05_manuscript_revised" / "Manuscript_working_stage3_sampling_density.docx",
    WORK_ROOT / "06_supplement_revised" / "Supplementary_materials_working_stage3_sampling_density.docx",
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def load_pilot_module():
    spec = importlib.util.spec_from_file_location("chronology_pilot", PILOT_SCRIPT)
    if spec is None or spec.loader is None:
        raise RuntimeError("Cannot import validated chronology pilot")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def read_expanded_membership(pilot, raw: pd.DataFrame) -> pd.DataFrame:
    wb = load_workbook(GROUPING_PATH, read_only=True, data_only=True)
    ws = wb["Expanded sensitivity"]
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    headers = [pilot.clean_id(x) for x in rows[0]]
    expanded = pd.DataFrame(rows[1:], columns=headers).dropna(how="all")
    expanded["operational_record"] = expanded.operational_record.map(pilot.clean_id)
    expanded["chronology_group_id"] = expanded.chronology_group_id.map(pilot.clean_id)
    expanded["n_measurements"] = expanded.n_measurements.astype(int)
    if len(expanded) != 2 or expanded.operational_record.nunique() != 2:
        raise RuntimeError("Expanded ODP 1210 membership must contain two records")
    if set(expanded.chronology_group_id) != {EXPANDED_GROUP_ID}:
        raise RuntimeError("Unexpected expanded chronology group")
    counts = raw.groupby("record_id").size()
    if int(counts.loc[expanded.operational_record].sum()) != 83 or int(expanded.n_measurements.sum()) != 83:
        raise RuntimeError("Expanded ODP 1210 membership must contain 83 retained measurements")
    return expanded


def draw_group_shifts(pilot) -> tuple[dict[str, np.ndarray], dict]:
    seed_sequence = np.random.SeedSequence(MASTER_SEED)
    children = seed_sequence.spawn(len(GROUP_ORDER))
    shifts, seed_info = {}, {}
    for gid, child in zip(GROUP_ORDER, children):
        spec = pilot.EXPECTED_GROUPS[gid]
        rng = np.random.default_rng(child)
        shifts[gid] = rng.uniform(spec["lo_kyr"], spec["hi_kyr"], size=ITERATIONS)
        seed_info[gid] = {
            "master_entropy": int(child.entropy),
            "spawn_key": list(child.spawn_key),
            "bit_generator": rng.bit_generator.__class__.__name__,
        }
    return shifts, seed_info


def critical_shift_keys(pilot, record_arrays, record_ids, shifts_ma: np.ndarray) -> tuple[np.ndarray, np.ndarray]:
    """Map whole-group shifts to intervals with identical 50-kyr assignments."""
    lower = pilot.BIN_SPEC["lower"]
    width = pilot.BIN_SPEC["width"]
    count = pilot.BIN_SPEC["count"]
    edges = lower + np.arange(count + 1) * width
    lo, hi = float(np.min(shifts_ma)), float(np.max(shifts_ma))
    critical = []
    for rid in record_ids:
        ages = record_arrays[rid][0]
        values = (edges[:, None] - ages[None, :]).ravel()
        values = values[(values >= lo) & (values <= hi)]
        critical.append(values)
    critical_values = np.unique(np.concatenate(critical)) if critical else np.array([], dtype=float)
    keys = np.searchsorted(critical_values, shifts_ma, side="right")
    return keys, critical_values


def cached_group_contributions(pilot, record_arrays, record_ids, shifts_kyr, label):
    shifts_ma = np.asarray(shifts_kyr, dtype=float) / 1000.0
    keys, critical_values = critical_shift_keys(pilot, record_arrays, record_ids, shifts_ma)
    unique_keys, inverse, frequencies = np.unique(keys, return_inverse=True, return_counts=True)
    unique_sums = np.empty((len(unique_keys), pilot.BIN_SPEC["count"]), dtype=float)
    unique_counts = np.empty((len(unique_keys), pilot.BIN_SPEC["count"]), dtype=np.int16)
    totals = Counter()
    for j, (key, frequency) in enumerate(zip(unique_keys, frequencies)):
        representative = int(np.flatnonzero(keys == key)[0])
        sums, counts, diag = pilot.record_contributions(record_arrays, record_ids, float(shifts_ma[representative]))
        unique_sums[j] = sums
        unique_counts[j] = counts
        totals.update({name: int(value) * int(frequency) for name, value in diag.items()})
        if (j + 1) % 500 == 0 or j + 1 == len(unique_keys):
            print(f"{label}: {j + 1}/{len(unique_keys)} unique bin-assignment states", flush=True)
    return unique_sums[inverse], unique_counts[inverse], {
        "unique_assignment_states": int(len(unique_keys)),
        "critical_shift_boundaries": int(len(critical_values)),
        "diagnostic_totals": dict(totals),
    }


def summarize_bins(ev, ages, observed, prefix):
    q = np.quantile(ev[:prefix], [0.025, 0.5, 0.975], axis=0)
    return pd.DataFrame({
        "iterations": prefix,
        "midpoint_Ma": ages,
        "observed_EV": observed,
        "simulation_median": q[1],
        "simulation_2.5th": q[0],
        "simulation_97.5th": q[2],
        "envelope_width": q[2] - q[0],
        "observed_minus_simulation_median": observed - q[1],
    })


def summarize_contrasts(contrasts, prefix):
    q = np.quantile(contrasts[:prefix], [0.025, 0.5, 0.975], axis=0)
    return pd.DataFrame([
        {
            "iterations": prefix,
            "contrast": name,
            "simulation_median": float(q[1, i]),
            "simulation_2.5th": float(q[0, i]),
            "simulation_97.5th": float(q[2, i]),
            "interval_width": float(q[2, i] - q[0, i]),
        }
        for i, name in enumerate(CONTRAST_NAMES)
    ])


def peak_midpoints(ev, ages):
    return ages[np.argmax(ev, axis=1)]


def peak_table(peak_ages, prefix, all_ages):
    counts = pd.Series(peak_ages[:prefix]).value_counts().to_dict()
    return pd.DataFrame([
        {
            "iterations": prefix,
            "peak_midpoint_Ma": float(age),
            "frequency": int(counts.get(float(age), 0)),
            "proportion": float(counts.get(float(age), 0) / prefix),
        }
        for age in all_ages
        if counts.get(float(age), 0) > 0
    ])


def convergence_tables(ev, contrasts, peak_ages, ages, observed):
    bin_prefixes = pd.concat([summarize_bins(ev, ages, observed, n) for n in PREFIXES], ignore_index=True)
    contrast_prefixes = pd.concat([summarize_contrasts(contrasts, n) for n in PREFIXES], ignore_index=True)
    peak_prefixes = pd.concat([peak_table(peak_ages, n, ages) for n in PREFIXES], ignore_index=True)

    final_bins = bin_prefixes[bin_prefixes.iterations.eq(10000)].set_index("midpoint_Ma")
    metric_rows = []
    for prefix in PREFIXES:
        current = bin_prefixes[bin_prefixes.iterations.eq(prefix)].set_index("midpoint_Ma")
        for statistic in ["simulation_median", "simulation_2.5th", "simulation_97.5th", "envelope_width"]:
            difference = (current[statistic] - final_bins[statistic]).abs()
            metric_rows.append({
                "iterations": prefix,
                "statistic": statistic,
                "median_absolute_difference_vs_10000": float(difference.median()),
                "maximum_absolute_difference_vs_10000": float(difference.max()),
                "midpoint_of_maximum_difference_Ma": float(difference.idxmax()),
            })
    convergence_metrics = pd.DataFrame(metric_rows)

    final_contrasts = contrast_prefixes[contrast_prefixes.iterations.eq(10000)].set_index("contrast")
    contrast_convergence_rows = []
    for prefix in PREFIXES:
        current = contrast_prefixes[contrast_prefixes.iterations.eq(prefix)].set_index("contrast")
        for contrast in CONTRAST_NAMES:
            for statistic in ["simulation_median", "simulation_2.5th", "simulation_97.5th", "interval_width"]:
                contrast_convergence_rows.append({
                    "iterations": prefix,
                    "contrast": contrast,
                    "statistic": statistic,
                    "absolute_difference_vs_10000": float(abs(current.loc[contrast, statistic] - final_contrasts.loc[contrast, statistic])),
                })
    contrast_convergence = pd.DataFrame(contrast_convergence_rows)

    final_counts = Counter(peak_ages)
    final_prop = np.array([final_counts.get(float(age), 0) / 10000 for age in ages])
    peak_convergence_rows = []
    for prefix in PREFIXES:
        counts = Counter(peak_ages[:prefix])
        prop = np.array([counts.get(float(age), 0) / prefix for age in ages])
        modal = max(counts, key=counts.get)
        peak_convergence_rows.append({
            "iterations": prefix,
            "modal_peak_midpoint_Ma": float(modal),
            "modal_peak_frequency": int(counts[modal]),
            "modal_peak_proportion": float(counts[modal] / prefix),
            "total_variation_distance_vs_10000": float(0.5 * np.sum(np.abs(prop - final_prop))),
            "maximum_absolute_probability_difference_vs_10000": float(np.max(np.abs(prop - final_prop))),
        })
    peak_convergence = pd.DataFrame(peak_convergence_rows)
    return bin_prefixes, convergence_metrics, contrast_prefixes, contrast_convergence, peak_prefixes, peak_convergence


def expanded_comparison(pilot, record_arrays, expanded, primary_sums, primary_counts, primary_ev, primary_contrasts, shifts, ages):
    expanded_records = expanded.operational_record.tolist()
    baseline_sums, baseline_counts, _ = pilot.record_contributions(record_arrays, expanded_records, 0.0)
    shifted_sums, shifted_counts, cache = cached_group_contributions(
        pilot, record_arrays, expanded_records,
        shifts["CHR-P1-ODP1209"][:EXPANDED_ITERATIONS],
        "Expanded ODP1210",
    )
    sums = primary_sums[:EXPANDED_ITERATIONS] - baseline_sums + shifted_sums
    counts = primary_counts[:EXPANDED_ITERATIONS] - baseline_counts + shifted_counts
    ev = sums / counts
    contrasts = pilot.phase_contrasts(ev, ages)
    peaks = peak_midpoints(ev, ages)
    primary_ev_1000 = primary_ev[:EXPANDED_ITERATIONS]
    primary_contrasts_1000 = primary_contrasts[:EXPANDED_ITERATIONS]
    primary_peaks = peak_midpoints(primary_ev_1000, ages)

    phase_rows = []
    for i, name in enumerate(CONTRAST_NAMES):
        pq = np.quantile(primary_contrasts_1000[:, i], [0.025, 0.5, 0.975])
        eq = np.quantile(contrasts[:, i], [0.025, 0.5, 0.975])
        dq = np.quantile(contrasts[:, i] - primary_contrasts_1000[:, i], [0.025, 0.5, 0.975])
        phase_rows.append({
            "contrast": name,
            "primary_median": float(pq[1]), "primary_2.5th": float(pq[0]), "primary_97.5th": float(pq[2]),
            "expanded_median": float(eq[1]), "expanded_2.5th": float(eq[0]), "expanded_97.5th": float(eq[2]),
            "paired_difference_median": float(dq[1]), "paired_difference_2.5th": float(dq[0]), "paired_difference_97.5th": float(dq[2]),
        })
    phase_comparison = pd.DataFrame(phase_rows)

    pcounts, ecounts = Counter(primary_peaks), Counter(peaks)
    peak_comparison = pd.DataFrame([
        {
            "peak_midpoint_Ma": float(age),
            "primary_frequency": int(pcounts.get(float(age), 0)),
            "primary_proportion": float(pcounts.get(float(age), 0) / EXPANDED_ITERATIONS),
            "expanded_frequency": int(ecounts.get(float(age), 0)),
            "expanded_proportion": float(ecounts.get(float(age), 0) / EXPANDED_ITERATIONS),
            "expanded_minus_primary_proportion": float((ecounts.get(float(age), 0) - pcounts.get(float(age), 0)) / EXPANDED_ITERATIONS),
        }
        for age in ages if pcounts.get(float(age), 0) or ecounts.get(float(age), 0)
    ])

    primary_q = np.quantile(primary_ev_1000, [0.025, 0.5, 0.975], axis=0)
    expanded_q = np.quantile(ev, [0.025, 0.5, 0.975], axis=0)
    bin_difference = pd.DataFrame({
        "midpoint_Ma": ages,
        "primary_median": primary_q[1],
        "expanded_median": expanded_q[1],
        "expanded_minus_primary_median": expanded_q[1] - primary_q[1],
        "primary_interval_width": primary_q[2] - primary_q[0],
        "expanded_interval_width": expanded_q[2] - expanded_q[0],
        "expanded_minus_primary_interval_width": (expanded_q[2] - expanded_q[0]) - (primary_q[2] - primary_q[0]),
    })
    max_bin = int(np.argmax(np.abs(bin_difference.expanded_minus_primary_median)))
    per_iteration_max = np.max(np.abs(ev - primary_ev_1000), axis=1)
    summary = {
        "analysis": "Expanded ODP 1210 chronology sensitivity pilot kept separate from primary",
        "iterations": EXPANDED_ITERATIONS,
        "additional_records": 2,
        "additional_measurements": 83,
        "shared_shift": "Same CHR-P1-ODP1209 Shatsky draw; no independent ODP1210 draw",
        "iterations_with_changed_peak_midpoint": int(np.sum(peaks != primary_peaks)),
        "proportion_with_changed_peak_midpoint": float(np.mean(peaks != primary_peaks)),
        "maximum_absolute_difference_between_bin_medians": float(abs(bin_difference.iloc[max_bin].expanded_minus_primary_median)),
        "maximum_difference_midpoint_Ma": float(bin_difference.iloc[max_bin].midpoint_Ma),
        "per_iteration_maximum_absolute_EV_difference_median": float(np.median(per_iteration_max)),
        "per_iteration_maximum_absolute_EV_difference_97.5th": float(np.quantile(per_iteration_max, 0.975)),
        "empty_bin_events": int(np.sum(counts == 0)),
        "cache": cache,
    }
    return ev, contrasts, peaks, counts, phase_comparison, peak_comparison, bin_difference, summary


def main():
    pilot = load_pilot_module()
    if pilot.MASTER_SEED != MASTER_SEED or pilot.ITERATIONS != 1000:
        raise RuntimeError("Validated pilot seed or size changed")
    protected_before = {str(p): sha256(p) for p in PROTECTED_PATHS if p.exists()}

    raw = pd.read_csv(RAW_PATH).reset_index(drop=True)
    raw["record_id"] = raw.record_id.map(pilot.clean_id)
    if len(raw) != 10881 or raw.record_id.nunique() != 70:
        raise RuntimeError("Retained baseline must contain 10,881 measurements and 70 records")
    membership, groups = pilot.load_grouping_audit()
    group_members = pilot.validate_grouping(raw, membership, groups)
    expanded = read_expanded_membership(pilot, raw)
    record_arrays = pilot.prepare_record_arrays(raw)
    ages, observed, observed_counts, s1_diff, observed_zero_bins = pilot.validate_observed_baseline(raw, record_arrays)

    active_records = set(membership.operational_record)
    fixed_records = [rid for rid in record_arrays if rid not in active_records]
    fixed_sums, fixed_counts, fixed_diag = pilot.record_contributions(record_arrays, fixed_records, 0.0)
    shifts, seed_info = draw_group_shifts(pilot)

    group_sums, group_counts, cache_info = {}, {}, {}
    for gid in GROUP_ORDER:
        sums, counts, cache = cached_group_contributions(
            pilot, record_arrays, group_members[gid], shifts[gid], gid
        )
        group_sums[gid], group_counts[gid], cache_info[gid] = sums, counts, cache

    total_sums = np.broadcast_to(fixed_sums, (ITERATIONS, pilot.BIN_SPEC["count"])).copy()
    total_counts = np.broadcast_to(fixed_counts, (ITERATIONS, pilot.BIN_SPEC["count"])).copy()
    for gid in GROUP_ORDER:
        total_sums += group_sums[gid]
        total_counts += group_counts[gid]
    if np.any(total_counts == 0):
        raise RuntimeError("Formal primary analysis produced an empty Global EV bin")
    ev = total_sums / total_counts
    contrasts = pilot.phase_contrasts(ev, ages)
    peaks = peak_midpoints(ev, ages)

    # The formal prefix must be exactly the validated pilot.
    if not PILOT_NPZ.exists():
        raise RuntimeError("Validated pilot realization file is missing")
    pilot_npz = np.load(PILOT_NPZ)
    prefix_ev_diff = float(np.max(np.abs(ev[:1000] - pilot_npz["primary_EV_50kyr"])))
    prefix_count_diff = int(np.max(np.abs(total_counts[:1000] - pilot_npz["primary_eligible_record_counts"])))
    prefix_contrast_diff = float(np.max(np.abs(contrasts[:1000] - pilot_npz["primary_phase_contrasts"])))
    prefix_shift_diff = max(
        float(np.max(np.abs(shifts["CHR-P1-ODP1209"][:1000] - pilot_npz["ODP1209_shift_kyr"]))),
        float(np.max(np.abs(shifts["CHR-P2-U1403"][:1000] - pilot_npz["U1403_shift_kyr"]))),
        float(np.max(np.abs(shifts["CHR-P3-SEYMOUR"][:1000] - pilot_npz["Seymour_shift_kyr"]))),
    )
    if prefix_ev_diff > 1e-12 or prefix_count_diff != 0 or prefix_contrast_diff > 1e-12 or prefix_shift_diff != 0:
        raise RuntimeError(
            f"Formal prefix does not reproduce pilot: EV={prefix_ev_diff}; counts={prefix_count_diff}; "
            f"contrasts={prefix_contrast_diff}; shifts={prefix_shift_diff}"
        )

    # Independent direct recalculation spot checks protect the caching optimization.
    direct_checks = {}
    for iteration in [0, 999, 4999, 9999]:
        sums = fixed_sums.copy()
        counts = fixed_counts.copy()
        for gid in GROUP_ORDER:
            s, c, _ = pilot.record_contributions(
                record_arrays, group_members[gid], float(shifts[gid][iteration] / 1000.0)
            )
            sums += s; counts += c
        direct_checks[str(iteration + 1)] = {
            "max_EV_difference": float(np.max(np.abs(sums / counts - ev[iteration]))),
            "max_record_count_difference": int(np.max(np.abs(counts - total_counts[iteration]))),
        }
        if direct_checks[str(iteration + 1)]["max_EV_difference"] > 1e-12 or direct_checks[str(iteration + 1)]["max_record_count_difference"] != 0:
            raise RuntimeError(f"Cached contribution check failed at iteration {iteration + 1}")

    bin_prefixes, convergence_metrics, contrast_prefixes, contrast_convergence, peak_prefixes, peak_convergence = convergence_tables(
        ev, contrasts, peaks, ages, observed
    )
    final_ev = bin_prefixes[bin_prefixes.iterations.eq(10000)].copy()
    final_contrasts = contrast_prefixes[contrast_prefixes.iterations.eq(10000)].copy()
    final_peaks = peak_prefixes[peak_prefixes.iterations.eq(10000)].copy()

    eligibility = pd.DataFrame({
        "midpoint_Ma": ages,
        "observed_eligible_record_bins": observed_counts,
        "simulation_record_bins_min": np.min(total_counts, axis=0),
        "simulation_record_bins_median": np.median(total_counts, axis=0),
        "simulation_record_bins_max": np.max(total_counts, axis=0),
        "simulation_record_bins_2.5th": np.quantile(total_counts, 0.025, axis=0),
        "simulation_record_bins_97.5th": np.quantile(total_counts, 0.975, axis=0),
        "iterations_with_empty_bin": np.sum(total_counts == 0, axis=0),
    })

    expanded_ev, expanded_contrasts, expanded_peaks, expanded_counts, expanded_phase, expanded_peak, expanded_bin_diff, expanded_summary = expanded_comparison(
        pilot, record_arrays, expanded, total_sums, total_counts, ev, contrasts, shifts, ages
    )

    # Reuse, do not rerun, the pilot group-contribution diagnostic.
    pilot_group_overview = pd.read_csv(PILOT_DIR / "Chronology_group_contribution_overview.csv")
    pilot_group_contrasts = pd.read_csv(PILOT_DIR / "Chronology_group_contribution_phase_contrasts.csv")
    pilot_group_overview.to_csv(RESULT_DIR / "Chronology_group_contribution_pilot_overview_reused.csv", index=False)
    pilot_group_contrasts.to_csv(RESULT_DIR / "Chronology_group_contribution_pilot_phase_contrasts_reused.csv", index=False)

    # Formal tabular outputs.
    final_ev.to_csv(RESULT_DIR / "Chronology_primary_formal_50kyr_Global_EV_summary.csv", index=False)
    final_contrasts.to_csv(RESULT_DIR / "Chronology_primary_formal_phase_contrasts.csv", index=False)
    final_peaks.to_csv(RESULT_DIR / "Chronology_primary_formal_peak_age_frequency.csv", index=False)
    eligibility.to_csv(RESULT_DIR / "Chronology_primary_formal_eligibility_diagnostics.csv", index=False)
    bin_prefixes.to_csv(RESULT_DIR / "Chronology_primary_formal_convergence_bin_prefixes.csv", index=False)
    convergence_metrics.to_csv(RESULT_DIR / "Chronology_primary_formal_convergence_metrics.csv", index=False)
    contrast_prefixes.to_csv(RESULT_DIR / "Chronology_primary_formal_convergence_phase_contrasts.csv", index=False)
    contrast_convergence.to_csv(RESULT_DIR / "Chronology_primary_formal_convergence_phase_metrics.csv", index=False)
    peak_prefixes.to_csv(RESULT_DIR / "Chronology_primary_formal_convergence_peak_distributions.csv", index=False)
    peak_convergence.to_csv(RESULT_DIR / "Chronology_primary_formal_convergence_peak_metrics.csv", index=False)
    expanded_phase.to_csv(RESULT_DIR / "Expanded_ODP1210_phase_contrast_comparison.csv", index=False)
    expanded_peak.to_csv(RESULT_DIR / "Expanded_ODP1210_peak_distribution_comparison.csv", index=False)
    expanded_bin_diff.to_csv(RESULT_DIR / "Expanded_ODP1210_primary_difference_by_bin.csv", index=False)
    (RESULT_DIR / "Expanded_ODP1210_summary.json").write_text(json.dumps(expanded_summary, indent=2, ensure_ascii=False), encoding="utf-8")

    shift_table = pd.DataFrame({
        "iteration": np.arange(1, ITERATIONS + 1),
        "CHR-P1-ODP1209_shift_kyr": shifts["CHR-P1-ODP1209"],
        "CHR-P2-U1403_shift_kyr": shifts["CHR-P2-U1403"],
        "CHR-P3-SEYMOUR_shift_kyr": shifts["CHR-P3-SEYMOUR"],
    })
    shift_table.to_csv(RESULT_DIR / "Chronology_primary_formal_group_shift_draws.csv", index=False)

    np.savez_compressed(
        RESULT_DIR / "Chronology_primary_formal_Global_EV_realizations_50kyr_10000.npz",
        age_50kyr=ages,
        observed_EV_50kyr=observed,
        primary_EV_50kyr=ev,
        primary_eligible_record_counts=total_counts,
        primary_phase_contrasts=contrasts,
        primary_peak_midpoint_Ma=peaks,
        ODP1209_shift_kyr=shifts["CHR-P1-ODP1209"],
        U1403_shift_kyr=shifts["CHR-P2-U1403"],
        Seymour_shift_kyr=shifts["CHR-P3-SEYMOUR"],
        expanded_ODP1210_EV_50kyr_1000=expanded_ev,
        expanded_ODP1210_eligible_record_counts_1000=expanded_counts,
        expanded_ODP1210_phase_contrasts_1000=expanded_contrasts,
        expanded_ODP1210_peak_midpoint_Ma_1000=expanded_peaks,
    )

    widths = final_ev.envelope_width.to_numpy(float)
    differences = np.abs(final_ev.observed_EV - final_ev.simulation_median).to_numpy(float)
    total_record_bins = total_counts.sum(axis=1)
    metadata = {
        "analysis": "Formal primary chronological sensitivity — standard 50-kyr Global EV",
        "iterations": ITERATIONS,
        "master_random_seed": MASTER_SEED,
        "group_random_streams": seed_info,
        "active_groups": 3,
        "active_operational_records": 24,
        "active_measurements": 4230,
        "active_measurement_coverage_percent": 100 * 4230 / 10881,
        "measurement_values_perturbed": False,
        "common_absolute_timescale_systematic_included": False,
        "ODP1210_in_primary": False,
        "dynamic_rebinning_and_eligibility": True,
        "Dataset_S1_observed_EV_max_absolute_validation_difference": s1_diff,
        "formal_prefix_vs_validated_pilot_max_EV_difference": prefix_ev_diff,
        "formal_prefix_vs_validated_pilot_max_record_count_difference": prefix_count_diff,
        "formal_prefix_vs_validated_pilot_max_contrast_difference": prefix_contrast_diff,
        "formal_prefix_vs_validated_pilot_max_shift_difference": prefix_shift_diff,
        "cache_info": cache_info,
        "direct_recalculation_checks": direct_checks,
        "observed_EV_vs_simulation_median_max_absolute_difference": float(np.max(differences)),
        "maximum_difference_midpoint_Ma": float(final_ev.iloc[int(np.argmax(differences))].midpoint_Ma),
        "envelope_median_width": float(np.median(widths)),
        "envelope_maximum_width": float(np.max(widths)),
        "maximum_width_midpoint_Ma": float(final_ev.iloc[int(np.argmax(widths))].midpoint_Ma),
        "observed_peak_midpoint_Ma": float(ages[int(np.argmax(observed))]),
        "empty_bin_events": int(np.sum(total_counts == 0)),
        "invalid_or_missing_EV_trajectories": int(np.sum(~np.all(np.isfinite(ev), axis=1))),
        "nonfinite_EV_values": int(np.sum(~np.isfinite(ev))),
        "observed_total_eligible_record_bins": int(observed_counts.sum()),
        "simulated_total_eligible_record_bins_min": int(total_record_bins.min()),
        "simulated_total_eligible_record_bins_median": float(np.median(total_record_bins)),
        "simulated_total_eligible_record_bins_max": int(total_record_bins.max()),
        "observed_zero_IQR_record_bins_excluded": int(observed_zero_bins),
        "expanded_ODP1210": expanded_summary,
        "group_contribution_diagnostic": "Pilot results reused without rerunning group-only simulations.",
    }
    (RESULT_DIR / "formal_metadata.json").write_text(json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8")

    protected_after = {str(p): sha256(p) for p in PROTECTED_PATHS if p.exists()}
    protected_check = pd.DataFrame([
        {"path": p, "sha256_before": h, "sha256_after": protected_after.get(p), "unchanged": h == protected_after.get(p)}
        for p, h in protected_before.items()
    ])
    protected_check.to_csv(RESULT_DIR / "protected_files_hash_check.csv", index=False)
    if not protected_check.unchanged.all():
        raise RuntimeError("Protected Dataset/document/source hash changed")

    manifest_paths = [RAW_PATH, IQR_PATH, S1_PATH, GROUPING_PATH, PILOT_SCRIPT, PILOT_NPZ, SCRIPT]
    pd.DataFrame([
        {"file": p.name, "path": str(p), "sha256": sha256(p)} for p in manifest_paths
    ]).to_csv(RESULT_DIR / "input_source_hash_manifest.csv", index=False)

    print(json.dumps(metadata, indent=2, ensure_ascii=False), flush=True)
    print("\nFORMAL PHASE CONTRASTS", flush=True)
    print(final_contrasts.to_string(index=False), flush=True)
    print("\nFORMAL PEAK DISTRIBUTION", flush=True)
    print(final_peaks.to_string(index=False), flush=True)
    print("\nCONVERGENCE METRICS", flush=True)
    print(convergence_metrics.to_string(index=False), flush=True)
    print("\nEXPANDED ODP1210 PHASE COMPARISON", flush=True)
    print(expanded_phase.to_string(index=False), flush=True)
    print("\nEXPANDED ODP1210 SUMMARY", flush=True)
    print(json.dumps(expanded_summary, indent=2, ensure_ascii=False), flush=True)


if __name__ == "__main__":
    main()
