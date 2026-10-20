#!/usr/bin/env python3
"""Layer 1 documented measurement-uncertainty pilot (1,000 iterations).

Only observation-level uncertainties that can be mapped one-to-one to retained
measurements are perturbed. Ages, bin assignments, record definitions, and the
observed eligible/non-zero record-bin mask are fixed.
"""

from __future__ import annotations

import hashlib
import json
import math
import shutil
import urllib.request
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


SEED = 20260823
ITERATIONS = 1000

SCRIPT = Path(__file__).resolve()
WORK_ROOT = SCRIPT.parent.parent
PROJECT_ROOT = WORK_ROOT  # package-relative provenance root
DATA_DIR = WORK_ROOT / "01_data_revised"
FORMAL_DIR = DATA_DIR / "formal_datasets"
RESULT_DIR = WORK_ROOT / "03_results" / "EV_uncertainty_inventory" / "Layer1_measurement_uncertainty_pilot_1000"
SOURCE_DIR = RESULT_DIR / "source_inputs"
RESULT_DIR.mkdir(parents=True, exist_ok=True)
SOURCE_DIR.mkdir(parents=True, exist_ok=True)

RAW_PATH = DATA_DIR / "retained_measurements_with_zscores_v4.csv"
IQR_PATH = DATA_DIR / "record_level_IQR_all_bin_widths_revised_v4_n2.csv"
S1_PATH = FORMAL_DIR / "Dataset_S1_revised_v6.xlsx"
S2_PATH = FORMAL_DIR / "Dataset_S2_revised.xlsx"
PARAM_PATH = WORK_ROOT / "03_results" / "EV_uncertainty_inventory" / "Executable_uncertainty_parameter_table.xlsx"
JOUINI_PATH = SOURCE_DIR / "Jouini_Appendix_C.xlsx"

PANGAEA_URLS = {
    "PANGAEA_707182.tab": "https://doi.pangaea.de/10.1594/PANGAEA.707182?format=textfile",
    "PANGAEA_785719.tab": "https://doi.pangaea.de/10.1594/PANGAEA.785719?format=textfile",
    "PANGAEA_785718.tab": "https://doi.pangaea.de/10.1594/PANGAEA.785718?format=textfile",
    "PANGAEA_785713.tab": "https://doi.pangaea.de/10.1594/PANGAEA.785713?format=textfile",
    "PANGAEA_785724.tab": "https://doi.pangaea.de/10.1594/PANGAEA.785724?format=textfile",
}

LAYER1_RECORDS = [
    "temperature::ODP1049",
    "temperature::Williston Basin",
    "temperature::Seymour Island",
    "temperature::Seymour Island 2",
    "temperature::Seymour Island 3",
    "d13C::ODP1209 5",
    "d18O::ODP1209 5",
    "weathering::Bottaccione Section, Gubbio Os",
    "weathering::ODP525A Os",
    "weathering::ODP577 Os",
    "weathering::ODP690C Os",
    "weathering::IODP U1403 Os",
    "weathering::ODP1209C Ca",
    "weathering::ODP1209C Ca2",
]

BIN_SPECS = {
    "50": {"width": 0.05, "lower": 64.775, "count": 49},
    "100": {"width": 0.10, "lower": 64.75, "count": 25},
    "150": {"width": 0.15, "lower": 63.975, "count": 27},
    "200": {"width": 0.20, "lower": 63.10, "count": 29},
}

PHASES = {
    "pre-escalation": (66.95, 67.20),
    "Phase 1": (66.45, 66.90),
    "Phase 2": (66.05, 66.40),
    "Recovery": (64.80, 65.95),
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def ensure_sources() -> dict[str, Path]:
    out: dict[str, Path] = {}
    for name, url in PANGAEA_URLS.items():
        p = SOURCE_DIR / name
        if not p.exists():
            raise FileNotFoundError(f"Bundled uncertainty source missing: {p}")
        out[name] = p
    for src, name in [
        (JOUINI_PATH, "Jouini_Appendix_C.xlsx"),
        (PARAM_PATH, "Executable_uncertainty_parameter_table.xlsx"),
    ]:
        dst = SOURCE_DIR / name
        if not dst.exists():
            shutil.copy2(src, dst)
        out[name] = dst
    return out


def read_pangaea(path: Path) -> pd.DataFrame:
    lines = path.read_text(encoding="utf-8").splitlines()
    marker = max(i for i, line in enumerate(lines) if line.strip() == "*/")
    return pd.read_csv(path, sep="\t", skiprows=marker + 1)


def numeric(x) -> bool:
    return isinstance(x, (int, float)) and math.isfinite(float(x))


@dataclass
class Mapping:
    row_index: int
    distribution: str
    sigma: float | None = None
    half_width: float | None = None
    source: str = ""
    source_key: str = ""
    isotope_value: float | None = None
    species: str | None = None


def validate_layer1_parameter_membership() -> None:
    wb = load_workbook(PARAM_PATH, read_only=True, data_only=True)
    ws = wb["Executable parameters"]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(x).strip() if x is not None else "" for x in rows[0]]
    idx = {h: i for i, h in enumerate(headers)}
    found = []
    for row in rows[1:]:
        if row[idx["measurement_analysis_layer"]] == "documented measurement propagation":
            found.append(str(row[idx["record_id"]]).strip())
    if set(found) != set(LAYER1_RECORDS) or len(found) != 14:
        raise RuntimeError(f"Layer 1 membership differs from the executable table: {found}")


def map_temperature_sources(raw: pd.DataFrame, sources: dict[str, Path]) -> tuple[list[Mapping], dict]:
    mappings: list[Mapping] = []
    notes = {}

    # ODP1049: the retained 76-temperature sequence maps exactly to the 29
    # S. beccariiformis + 47 N. truempyi observations in PANGAEA 707182.
    src = read_pangaea(sources["PANGAEA_707182.tab"])
    s_o = next(c for c in src if "S. beccariiformis δ18O [" in c)
    n_o = next(c for c in src if "N. truempyi δ18O [" in c)
    s_e = next(c for c in src if "δ18O std dev" in c and "S. beccariiformis" in c)
    n_e = next(c for c in src if "δ18O std dev" in c and "N. truempyi" in c)
    obs = []
    for source_row, row in src.iterrows():
        for species, value_col, error_col in [("S", s_o, s_e), ("N", n_o, n_e)]:
            if pd.notna(row[value_col]):
                obs.append((source_row, species, float(row[value_col]), row[error_col]))
    retained = raw[raw.record_id == "temperature::ODP1049"].sort_values("age_Ma")
    if len(retained) != len(obs):
        raise RuntimeError("ODP1049 retained/source sequence lengths differ")
    for (source_row, species, isotope, error), (row_index, row) in zip(obs, retained.iterrows()):
        predicted = (0.03 * isotope**2 - 4.43 * isotope + 10.2855) if species == "S" else (0.03 * isotope**2 - 4.427 * isotope + 10.064075)
        if abs(predicted - float(row.value)) > 1e-8:
            raise RuntimeError("ODP1049 isotope-to-temperature mapping failed")
        if numeric(error):
            mappings.append(Mapping(row_index, "odp1049_isotope_normal", sigma=float(error), source="PANGAEA 707182 δ18O SD", source_key=f"row={source_row + 1};species={species}", isotope_value=isotope, species=species))
    notes["temperature::ODP1049"] = f"{sum(m.row_index in retained.index for m in mappings)} of {len(retained)} observations have numeric source δ18O SD; the remainder stay fixed."

    wb = load_workbook(S2_PATH, read_only=True, data_only=True)
    direct = [
        ("temperature::Williston Basin", "Williston Basin"),
        ("temperature::Seymour Island", "Seymour Island"),
        ("temperature::Seymour Island 2", "Seymour Island 2"),
        ("temperature::Seymour Island 3", "Seymour Island 3"),
    ]
    for rid, sheet in direct:
        ws = wb[sheet]
        retained = raw[raw.record_id == rid]
        count = 0
        for row_index, row in retained.iterrows():
            source_row = int(row.source_row)
            age = ws.cell(source_row, 3).value
            value = ws.cell(source_row, 6).value
            error = ws.cell(source_row, 7).value
            if not (numeric(age) and numeric(value)) or abs(float(age) - float(row.age_Ma)) > 1e-7 or abs(float(value) - float(row.value)) > 1e-7:
                raise RuntimeError(f"{rid} Dataset S2 row mapping failed at source row {source_row}")
            if numeric(error):
                mappings.append(Mapping(row_index, "bounded_uniform", half_width=abs(float(error)), source=f"Dataset S2/{sheet} reported ± field", source_key=f"sheet={sheet};row={source_row}"))
                count += 1
        notes[rid] = f"{count} of {len(retained)} observations have numeric source ± values; nonnumeric/NA fields stay fixed."
    return mappings, notes


def unique_age_value_match(source_rows, retained: pd.DataFrame, age_col: int, value_col: int, error_col: int, age_tol=0.001, value_tol=1e-9):
    mapped = []
    candidates = [r for r in source_rows if numeric(r[age_col]) and numeric(r[value_col])]
    for row_index, row in retained.iterrows():
        hits = [r for r in candidates if abs(float(r[age_col]) - float(row.age_Ma)) <= age_tol and abs(float(r[value_col]) - float(row.value)) <= value_tol]
        if len(hits) != 1:
            raise RuntimeError(f"Expected one source match for {row.record_id} at {row.age_Ma}/{row.value}; found {len(hits)}")
        mapped.append((row_index, hits[0]))
    return mapped


def map_jouini_sources(raw: pd.DataFrame, sources: dict[str, Path]) -> tuple[list[Mapping], dict]:
    wb = load_workbook(sources["Jouini_Appendix_C.xlsx"], read_only=True, data_only=True)
    p_rows = list(wb["Appendic_C.1"].iter_rows(min_row=5, values_only=True))
    b_rows = list(wb["Appendic_C.2"].iter_rows(min_row=4, values_only=True))
    specs = [
        ("d13C::ODP1209 5", p_rows, 4, 21, 23, "Appendic_C.1 δ13C ±2SD"),
        ("d18O::ODP1209 5", p_rows, 4, 22, 24, "Appendic_C.1 δ18O ±2SD"),
        ("weathering::ODP1209C Ca", p_rows, 4, 7, 9, "Appendic_C.1 δ44/40CaP ±2SD"),
        ("weathering::ODP1209C Ca2", b_rows, 4, 6, 7, "Appendic_C.2 δ44/40CaB ±2SD"),
    ]
    mappings, notes = [], {}
    for rid, rows, age_col, value_col, error_col, label in specs:
        retained = raw[raw.record_id == rid]
        matched = unique_age_value_match(rows, retained, age_col, value_col, error_col)
        count = 0
        for row_index, source_row in matched:
            error = source_row[error_col]
            if numeric(error):
                mappings.append(Mapping(row_index, "normal", sigma=abs(float(error)) / 2.0, source=label, source_key=f"age={float(source_row[age_col]):.9f};value={float(source_row[value_col]):.12g}"))
                count += 1
        notes[rid] = f"{count} of {len(retained)} observations have mapped numeric ±2SD; missing source errors stay fixed."
    return mappings, notes


def greedy_subsequence_map(retained: pd.DataFrame, source: pd.DataFrame, value_col: str, error_col: str, start_retained: int = 0):
    source = source.dropna(subset=[value_col]).reset_index(drop=True)
    mapped, pos = [], -1
    for row_index, row in retained.iloc[start_retained:].iterrows():
        candidates = [j for j in range(pos + 1, len(source)) if abs(float(source.loc[j, value_col]) - float(row.value)) <= 1e-9]
        if not candidates:
            return None
        pos = candidates[0]
        mapped.append((row_index, source.loc[pos]))
    return mapped


def map_os_sources(raw: pd.DataFrame, sources: dict[str, Path]) -> tuple[list[Mapping], dict]:
    mappings, notes = [], {}

    gubbio = read_pangaea(sources["PANGAEA_785719.tab"])
    retained = raw[raw.record_id == "weathering::Bottaccione Section, Gubbio Os"].sort_values("age_Ma")
    matched = greedy_subsequence_map(retained, gubbio, "187Os/188Os", "187Os/188Os e [±]")
    if matched is None or len(matched) != len(retained):
        raise RuntimeError("Gubbio PANGAEA sequence mapping failed")
    for row_index, source_row in matched:
        mappings.append(Mapping(row_index, "normal", sigma=abs(float(source_row["187Os/188Os e [±]"])) / 2.0, source="PANGAEA 785719 187Os/188Os 2σ", source_key=f"height={source_row['Height [m]']};value={source_row['187Os/188Os']}"))
    notes["weathering::Bottaccione Section, Gubbio Os"] = f"{len(matched)} of {len(retained)} observations mapped one-to-one to PANGAEA 785719."

    # ODP690C is mixed-source. The retained suffix beginning at index 27 maps
    # sequentially and exactly to PANGAEA 785718; earlier rows remain fixed.
    odp690 = read_pangaea(sources["PANGAEA_785718.tab"])
    retained = raw[raw.record_id == "weathering::ODP690C Os"].sort_values("age_Ma")
    matched = None
    start = None
    for candidate_start in range(len(retained)):
        candidate = greedy_subsequence_map(retained, odp690, "187Os/188Os", "187Os/188Os e [±]", candidate_start)
        if candidate is not None:
            matched, start = candidate, candidate_start
            break
    if matched is None or len(matched) < 2:
        raise RuntimeError("ODP690C PANGAEA suffix mapping failed")
    for row_index, source_row in matched:
        mappings.append(Mapping(row_index, "normal", sigma=abs(float(source_row["187Os/188Os e [±]"])) / 2.0, source="PANGAEA 785718 187Os/188Os 2σ", source_key=f"sample={source_row['Sample label']};value={source_row['187Os/188Os']}"))
    notes["weathering::ODP690C Os"] = f"{len(matched)} of {len(retained)} retained observations (ordered suffix starting at retained position {start + 1}) map one-to-one to PANGAEA 785718; earlier mixed-source rows stay fixed."

    notes["weathering::ODP525A Os"] = "0 mapped: retained values are not one-to-one with the available PANGAEA measurement/error rows; record stays fixed in this strict Layer 1 pilot."
    notes["weathering::ODP577 Os"] = "0 mapped: retained values are not one-to-one with the available PANGAEA measurement/error rows; record stays fixed in this strict Layer 1 pilot."
    notes["weathering::IODP U1403 Os"] = "0 mapped: the source reports an approximate record-level ±2.5% and references Table S16, but the retained rows lack accessible one-to-one 95% CI fields; record stays fixed."
    return mappings, notes


def assign_bins(ages: np.ndarray, spec: dict) -> np.ndarray:
    width, lower, count = spec["width"], spec["lower"], spec["count"]
    # Literal reproduction of Code 1/EV reconstruction's sequential
    # [lower, upper) loop, including its floating-point edge behavior.
    idx = np.full(len(ages), -1, dtype=int)
    for j, age in enumerate(ages):
        for i in range(count):
            lo = lower + i * width
            hi = lower + (i + 1) * width
            if lo <= age < hi:
                idx[j] = i
                break
    return idx


def iqr_linear(x: np.ndarray) -> float:
    y = np.sort(x)
    n = len(y)
    def q(p):
        h = (n - 1) * p
        lo = int(math.floor(h))
        hi = int(math.ceil(h))
        return y[lo] if lo == hi else y[lo] + (h - lo) * (y[hi] - y[lo])
    return float(q(0.75) - q(0.25))


def build_record_bin_masks(raw: pd.DataFrame, iqr_table: pd.DataFrame, record_code: np.ndarray):
    ages = raw.age_Ma.to_numpy(float)
    masks = {}
    baseline_diffs = []
    values = raw.value.to_numpy(float)
    z_obs = np.empty(len(raw))
    for code in np.unique(record_code):
        ix = np.flatnonzero(record_code == code)
        z_obs[ix] = (values[ix] - values[ix].mean()) / values[ix].std(ddof=1)

    for scheme, spec in BIN_SPECS.items():
        assigned = assign_bins(ages, spec)
        centers = spec["lower"] + (np.arange(spec["count"]) + 0.5) * spec["width"]
        mask_rows = iqr_table[(iqr_table.scheme.astype(str) == scheme) & (iqr_table.alignment == "standard") & (iqr_table.IQR > 0)].copy()
        groups = []
        bin_counts = np.zeros(spec["count"], dtype=int)
        observed_sums = np.zeros(spec["count"], dtype=float)
        record_lookup = {rid: i for i, rid in enumerate(pd.unique(raw.record_id))}
        for _, row in mask_rows.iterrows():
            rid = row.record_id
            center = float(row.bin_center_Ma)
            b = int(np.argmin(abs(centers - center)))
            code = record_lookup[rid]
            obs_ix = np.flatnonzero((record_code == code) & (assigned == b))
            if len(obs_ix) != int(row.n_measurements):
                raise RuntimeError(f"Mask n mismatch: {scheme}/{rid}/{center}: {len(obs_ix)} vs {row.n_measurements}")
            calc = iqr_linear(z_obs[obs_ix])
            if abs(calc - float(row.IQR)) > 5e-10:
                baseline_diffs.append((scheme, rid, center, calc, float(row.IQR)))
            groups.append((b, obs_ix))
            bin_counts[b] += 1
            observed_sums[b] += calc
        if np.any(bin_counts == 0):
            raise RuntimeError(f"Observed mask leaves empty bins in {scheme}")
        masks[scheme] = {"centers": centers, "groups": groups, "bin_counts": bin_counts, "observed_ev": observed_sums / bin_counts}
    if baseline_diffs:
        raise RuntimeError(f"Bottom-up baseline differs from record-IQR table; first differences: {baseline_diffs[:5]}")
    return masks


def validate_s1_50(masks) -> float:
    wb = load_workbook(S1_PATH, read_only=True, data_only=True)
    ws = wb["total IQR"]
    s1 = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) >= 16 and numeric(row[14]) and numeric(row[15]):
            s1[float(row[15]) / 1000.0] = float(row[14])
    centers, ev = masks["50"]["centers"], masks["50"]["observed_ev"]
    return max(abs(s1[round(float(age), 8)] - float(value)) for age, value in zip(centers, ev))


def simulate(raw: pd.DataFrame, mappings: list[Mapping], masks):
    rng = np.random.default_rng(SEED)
    record_ids = pd.unique(raw.record_id)
    record_lookup = {rid: i for i, rid in enumerate(record_ids)}
    record_code = raw.record_id.map(record_lookup).to_numpy(int)
    record_indices = [np.flatnonzero(record_code == code) for code in range(len(record_ids))]
    base_values = raw.value.to_numpy(float)

    normal = [m for m in mappings if m.distribution == "normal"]
    bounded = [m for m in mappings if m.distribution == "bounded_uniform"]
    isotope = [m for m in mappings if m.distribution == "odp1049_isotope_normal"]
    normal_ix = np.array([m.row_index for m in normal], dtype=int)
    normal_sigma = np.array([m.sigma for m in normal], dtype=float)
    bounded_ix = np.array([m.row_index for m in bounded], dtype=int)
    bounded_u = np.array([m.half_width for m in bounded], dtype=float)
    isotope_ix = np.array([m.row_index for m in isotope], dtype=int)
    isotope_sigma = np.array([m.sigma for m in isotope], dtype=float)
    isotope_base = np.array([m.isotope_value for m in isotope], dtype=float)
    isotope_species = np.array([m.species for m in isotope])

    outputs = {scheme: np.empty((ITERATIONS, len(info["centers"])), dtype=float) for scheme, info in masks.items()}
    contrast_names = ["Phase 1 - pre-escalation", "Phase 2 - pre-escalation", "Recovery - Phase 2"]
    contrasts = np.empty((ITERATIONS, 3), dtype=float)
    nonfinite_values = nonfinite_z = nonfinite_ev = zero_sd = 0
    global_min, global_max, max_abs_z = math.inf, -math.inf, 0.0
    perturbed_min = np.full(len(mappings), math.inf)
    perturbed_max = np.full(len(mappings), -math.inf)

    for it in range(ITERATIONS):
        values = base_values.copy()
        if len(normal_ix):
            values[normal_ix] += rng.normal(0.0, normal_sigma)
        if len(bounded_ix):
            values[bounded_ix] += rng.uniform(-bounded_u, bounded_u)
        if len(isotope_ix):
            iso = isotope_base + rng.normal(0.0, isotope_sigma)
            s = isotope_species == "S"
            values[isotope_ix[s]] = 0.03 * iso[s] ** 2 - 4.43 * iso[s] + 10.2855
            values[isotope_ix[~s]] = 0.03 * iso[~s] ** 2 - 4.427 * iso[~s] + 10.064075
        if not np.all(np.isfinite(values)):
            nonfinite_values += 1

        z = np.empty(len(raw), dtype=float)
        for ix in record_indices:
            sd = values[ix].std(ddof=1)
            if not math.isfinite(sd) or sd <= 0:
                zero_sd += 1
                z[ix] = np.nan
            else:
                z[ix] = (values[ix] - values[ix].mean()) / sd
        if not np.all(np.isfinite(z)):
            nonfinite_z += 1
        max_abs_z = max(max_abs_z, float(np.nanmax(abs(z))))

        for scheme, info in masks.items():
            sums = np.zeros(len(info["centers"]), dtype=float)
            for b, obs_ix in info["groups"]:
                sums[b] += iqr_linear(z[obs_ix])
            ev = sums / info["bin_counts"]
            outputs[scheme][it] = ev
            if not np.all(np.isfinite(ev)):
                nonfinite_ev += 1
            global_min = min(global_min, float(np.nanmin(ev)))
            global_max = max(global_max, float(np.nanmax(ev)))

        ages = masks["50"]["centers"]
        ev50 = outputs["50"][it]
        phase_mean = {}
        for name, (lo, hi) in PHASES.items():
            sel = (ages >= lo - 1e-10) & (ages <= hi + 1e-10)
            phase_mean[name] = float(ev50[sel].mean())
        contrasts[it] = [
            phase_mean["Phase 1"] - phase_mean["pre-escalation"],
            phase_mean["Phase 2"] - phase_mean["pre-escalation"],
            phase_mean["Recovery"] - phase_mean["Phase 2"],
        ]

    anomalies = {
        "nonfinite_value_iterations": nonfinite_values,
        "nonfinite_z_iterations": nonfinite_z,
        "nonfinite_EV_scheme_iterations": nonfinite_ev,
        "zero_or_invalid_record_SD_events": zero_sd,
        "maximum_absolute_simulated_z_score": max_abs_z,
        "minimum_simulated_global_EV_all_widths": global_min,
        "maximum_simulated_global_EV_all_widths": global_max,
    }
    return outputs, contrasts, contrast_names, anomalies


def summarize(raw, mappings, notes, masks, outputs, contrasts, contrast_names, anomalies, sources, s1_diff):
    summaries = []
    for scheme, arr in outputs.items():
        info = masks[scheme]
        q = np.quantile(arr, [0.025, 0.5, 0.975], axis=0)
        for age, observed, lo, med, hi in zip(info["centers"], info["observed_ev"], q[0], q[1], q[2]):
            summaries.append({"bin_width_kyr": int(scheme), "midpoint_Ma": age, "observed_EV": observed, "MC_median": med, "MC_2.5th": lo, "MC_97.5th": hi, "interval_width": hi - lo, "observed_minus_MC_median": observed - med})
    summary_df = pd.DataFrame(summaries).sort_values(["bin_width_kyr", "midpoint_Ma"], ascending=[True, False])
    summary_df.to_csv(RESULT_DIR / "Layer1_MC_Global_EV_summary_all_widths.csv", index=False)
    summary50 = summary_df[summary_df.bin_width_kyr == 50].copy()
    summary50.to_csv(RESULT_DIR / "Layer1_MC_Global_EV_50kyr_summary.csv", index=False)

    cq = np.quantile(contrasts, [0.025, 0.5, 0.975], axis=0)
    contrast_df = pd.DataFrame({"contrast": contrast_names, "MC_median": cq[1], "MC_2.5th": cq[0], "MC_97.5th": cq[2], "interval_width": cq[2] - cq[0]})
    contrast_df.to_csv(RESULT_DIR / "Layer1_MC_phase_contrasts_50kyr.csv", index=False)

    mapped_by_record = pd.Series([raw.loc[m.row_index, "record_id"] for m in mappings]).value_counts().to_dict()
    distribution_by_record = {}
    source_by_record = {}
    for m in mappings:
        rid = raw.loc[m.row_index, "record_id"]
        distribution_by_record.setdefault(rid, set()).add(m.distribution)
        source_by_record.setdefault(rid, set()).add(m.source)
    audit_rows = []
    for rid in LAYER1_RECORDS:
        retained_n = int((raw.record_id == rid).sum())
        mapped_n = int(mapped_by_record.get(rid, 0))
        audit_rows.append({
            "record_id": rid,
            "retained_measurements": retained_n,
            "perturbed_measurements": mapped_n,
            "coverage_percent": 100 * mapped_n / retained_n if retained_n else 0,
            "pilot_status": "perturbed" if mapped_n else "fixed_due_to_unresolved_row_mapping",
            "distribution": "; ".join(sorted(distribution_by_record.get(rid, {"fixed"}))),
            "source_uncertainty": "; ".join(sorted(source_by_record.get(rid, {"not executable at observation level"}))),
            "mapping_note": notes.get(rid, ""),
        })
    audit_df = pd.DataFrame(audit_rows)
    audit_df.to_csv(RESULT_DIR / "Layer1_observation_mapping_audit.csv", index=False)

    perturbed_records = int((audit_df.perturbed_measurements > 0).sum())
    perturbed_measurements = int(audit_df.perturbed_measurements.sum())
    widths = summary50.interval_width.to_numpy()
    max_abs = float(abs(summary50.observed_minus_MC_median).max())
    max_abs_row = summary50.iloc[int(abs(summary50.observed_minus_MC_median).argmax())]
    metadata = {
        "analysis": "Layer 1 documented measurement uncertainty propagation pilot",
        "iterations": ITERATIONS,
        "random_seed": SEED,
        "chronology_perturbed": False,
        "method_based_conservative_uncertainty_included": False,
        "retained_measurements": int(len(raw)),
        "layer1_parameter_records": 14,
        "actually_perturbed_records": perturbed_records,
        "actually_perturbed_measurements": perturbed_measurements,
        "perturbed_measurement_percent": 100 * perturbed_measurements / len(raw),
        "observed_baseline_nonzero_mask": True,
        "s1_50kyr_max_absolute_validation_difference": s1_diff,
        "observed_EV_vs_MC_median_max_absolute_difference": max_abs,
        "max_difference_midpoint_Ma": float(max_abs_row.midpoint_Ma),
        "50kyr_interval_width_median": float(np.median(widths)),
        "50kyr_interval_width_maximum": float(np.max(widths)),
        "50kyr_maximum_width_midpoint_Ma": float(summary50.iloc[int(widths.argmax())].midpoint_Ma),
        "anomaly_diagnostics": anomalies,
        "notes": [
            "Only retained observations with one-to-one mapped documented quantitative uncertainty were perturbed.",
            "ODP525A Os, ODP577 Os, and IODP U1403 Os remained fixed because the executable table's required observation-level source-key check could not be closed with available source rows.",
            "Ages and bin assignments were fixed; no chronological uncertainty was applied.",
            "The observed non-zero eligible record-bin mask was held fixed in every realization.",
        ],
    }
    (RESULT_DIR / "pilot_metadata.json").write_text(json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8")

    manifest = []
    for name, path in sources.items():
        manifest.append({"file": name, "sha256": sha256(path), "source": PANGAEA_URLS.get(name, str(path))})
    for path in [RAW_PATH, IQR_PATH, S1_PATH, S2_PATH, SCRIPT]:
        manifest.append({"file": path.name, "sha256": sha256(path), "source": str(path)})
    pd.DataFrame(manifest).to_csv(RESULT_DIR / "input_source_hash_manifest.csv", index=False)

    np.savez_compressed(
        RESULT_DIR / "Layer1_MC_Global_EV_realizations_all_widths.npz",
        **{f"EV_{scheme}kyr": arr for scheme, arr in outputs.items()},
        **{f"age_{scheme}kyr": masks[scheme]["centers"] for scheme in masks},
        phase_contrasts_50kyr=contrasts,
    )
    return metadata, summary50, contrast_df, audit_df


def main():
    validate_layer1_parameter_membership()
    sources = ensure_sources()
    raw = pd.read_csv(RAW_PATH).reset_index(drop=True)
    if len(raw) != 10881 or raw.record_id.nunique() != 70:
        raise RuntimeError(f"Unexpected retained baseline dimensions: {len(raw)} measurements, {raw.record_id.nunique()} records")
    iqr_table = pd.read_csv(IQR_PATH)

    mappings = []
    notes = {}
    for mapper in (map_temperature_sources, map_jouini_sources, map_os_sources):
        mm, nn = mapper(raw, sources)
        mappings.extend(mm)
        notes.update(nn)
    row_ids = [m.row_index for m in mappings]
    if len(row_ids) != len(set(row_ids)):
        raise RuntimeError("An observation was assigned more than one uncertainty mapping")

    record_ids = pd.unique(raw.record_id)
    record_lookup = {rid: i for i, rid in enumerate(record_ids)}
    record_code = raw.record_id.map(record_lookup).to_numpy(int)
    masks = build_record_bin_masks(raw, iqr_table, record_code)
    s1_diff = validate_s1_50(masks)
    if s1_diff > 1e-12:
        raise RuntimeError(f"Recomputed observed 50-kyr EV does not match Dataset S1 v6: {s1_diff}")

    outputs, contrasts, contrast_names, anomalies = simulate(raw, mappings, masks)
    metadata, summary50, contrast_df, audit_df = summarize(raw, mappings, notes, masks, outputs, contrasts, contrast_names, anomalies, sources, s1_diff)
    print(json.dumps(metadata, indent=2, ensure_ascii=False))
    print("\nPHASE CONTRASTS")
    print(contrast_df.to_string(index=False))
    print("\nMAPPING AUDIT")
    print(audit_df[["record_id", "retained_measurements", "perturbed_measurements", "pilot_status"]].to_string(index=False))


if __name__ == "__main__":
    main()
