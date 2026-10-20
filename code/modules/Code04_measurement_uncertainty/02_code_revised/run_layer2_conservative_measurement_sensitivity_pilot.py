#!/usr/bin/env python3
"""Layer 2 conservative method-based measurement-sensitivity pilot (50 kyr).

This 1,000-iteration pilot perturbs measurement values only. Chronology,
record definitions, bin assignments, and the observed non-zero record-bin
eligibility mask remain fixed. Only the standard 50-kyr primary Global EV
reconstruction is evaluated. Documented Layer 1A mappings are retained;
Tier 2 observation-level parameters are added exactly by their locked primary
distribution. Calibration-only rows that explicitly require unavailable raw
calibration variables/parameter covariance remain fixed rather than being
implemented as a constant shift that would cancel after z-standardization.
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import math
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path

import numpy as np
import pandas as pd


SEED = 20260824
ITERATIONS = 1000
CARBONATE_OR_PHOSPHATE_DT_PER_D18O = -5.0  # deg C per per mil; locked method-literature equivalence

SCRIPT = Path(__file__).resolve()
WORK_ROOT = SCRIPT.parent.parent
DATA_DIR = WORK_ROOT / "01_data_revised"
FORMAL_DIR = DATA_DIR / "formal_datasets"
RESULT_DIR = WORK_ROOT / "03_results" / "EV_uncertainty_inventory" / "Layer2_conservative_measurement_sensitivity_pilot_50kyr_1000"
RESULT_DIR.mkdir(parents=True, exist_ok=True)

RAW_PATH = DATA_DIR / "retained_measurements_with_zscores_v4.csv"
IQR_PATH = DATA_DIR / "record_level_IQR_all_bin_widths_revised_v4_n2.csv"
S1_PATH = FORMAL_DIR / "Dataset_S1_revised_v6.xlsx"
PARAM_PATH = WORK_ROOT / "03_results" / "EV_uncertainty_inventory" / "Executable_uncertainty_parameter_table.xlsx"
LAYER1_SCRIPT = WORK_ROOT / "02_code_revised" / "run_layer1_measurement_uncertainty_pilot.py"

PROTECTED_PATHS = [
    FORMAL_DIR / "Dataset_S1_revised_v6.xlsx",
    FORMAL_DIR / "Dataset_S2_revised.xlsx",
    FORMAL_DIR / "Dataset_S3_revised_v3.xlsx",
    FORMAL_DIR / "Dataset_S4_revised_v3.xlsx",
    FORMAL_DIR / "Dataset_S5_revised_v4.xlsx",
    FORMAL_DIR / "Dataset_S6_revised_v4.xlsx",
    WORK_ROOT / "05_manuscript_revised" / "Manuscript_working_stage3_sampling_density.docx",
    WORK_ROOT / "06_supplement_revised" / "Supplementary_materials_working_stage3_sampling_density.docx",
]

BIN_SPECS = {
    "50": {"width": 0.05, "lower": 64.775, "count": 49},
}

PHASES = {
    "pre-escalation": (66.95, 67.20),
    "Phase 1": (66.45, 66.90),
    "Phase 2": (66.05, 66.40),
    "Recovery": (64.80, 65.95),
}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def load_layer1_module():
    spec = importlib.util.spec_from_file_location("layer1_pilot", LAYER1_SCRIPT)
    if spec is None or spec.loader is None:
        raise RuntimeError("Could not load the preserved Layer 1 pilot code")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


@dataclass
class DrawMap:
    row_index: int
    distribution: str
    sigma: float | None = None
    lower: float | None = None
    upper: float | None = None
    isotope_value: float | None = None
    species: str | None = None
    implementation: str = ""
    source_layer: str = ""


def build_layer1_maps(raw: pd.DataFrame):
    l1 = load_layer1_module()
    l1.validate_layer1_parameter_membership()
    sources = l1.ensure_sources()
    mappings = []
    notes = {}
    for mapper in (l1.map_temperature_sources, l1.map_jouini_sources, l1.map_os_sources):
        mm, nn = mapper(raw, sources)
        notes.update(nn)
        for m in mm:
            if m.distribution == "normal":
                mappings.append(DrawMap(m.row_index, "normal_direct", sigma=float(m.sigma), implementation="documented observation-level Gaussian", source_layer="Layer 1A documented"))
            elif m.distribution == "bounded_uniform":
                mappings.append(DrawMap(m.row_index, "uniform_direct", lower=-float(m.half_width), upper=float(m.half_width), implementation="documented observation-level bounded uniform", source_layer="Layer 1A documented"))
            elif m.distribution == "odp1049_isotope_normal":
                mappings.append(DrawMap(m.row_index, "odp1049_isotope_normal", sigma=float(m.sigma), isotope_value=float(m.isotope_value), species=m.species, implementation="documented δ18O Gaussian followed by exact retained ODP1049 reconstruction", source_layer="Layer 1A documented"))
            else:
                raise RuntimeError(f"Unsupported Layer 1 mapping distribution: {m.distribution}")
    return mappings, notes, sources


def parameter_table() -> pd.DataFrame:
    p = pd.read_excel(PARAM_PATH, sheet_name="Executable parameters")
    p["record_id"] = p.record_id.map(lambda x: unicodedata.normalize("NFC", str(x)).strip())
    if len(p) != 70 or p.record_id.nunique() != 70:
        raise RuntimeError("Executable table does not contain exactly 70 unique operational records")
    return p


def tier2_maps(raw: pd.DataFrame, params: pd.DataFrame):
    maps: list[DrawMap] = []
    audit = []
    tier2 = params[params.measurement_analysis_layer.eq("conservative measurement sensitivity")]
    if len(tier2) != 56:
        raise RuntimeError(f"Expected 56 Tier 2 records, found {len(tier2)}")

    for _, r in tier2.iterrows():
        rid = r.record_id
        rows = raw.index[raw.record_id.eq(rid)].to_numpy(int)
        if len(rows) != int(r.n_measurements):
            raise RuntimeError(f"Parameter/raw count mismatch for {rid}: {len(rows)} vs {r.n_measurements}")
        dist = r.measurement_distribution
        target = str(r.measurement_perturbation_target)
        status = str(r.measurement_execution_status)
        implementation = ""
        active = False

        if dist == "normal_observation_level":
            sigma = float(r.measurement_sigma_max)
            if "Underlying" in target:
                sigma *= abs(CARBONATE_OR_PHOSPHATE_DT_PER_D18O)
                implementation = "Gaussian δ18O uncertainty converted by locked method-literature local slope −5 °C/‰ before record-wide restandardization"
            else:
                implementation = "Gaussian observation-level perturbation in retained measurement units"
            maps.extend(DrawMap(i, "normal_direct", sigma=sigma, implementation=implementation, source_layer="Layer 2 conservative Gaussian") for i in rows)
            active = True

        elif dist == "bounded_uniform_observation_level":
            lower = float(r.measurement_bound_lower)
            upper = float(r.measurement_bound_upper)
            if "Underlying" in target:
                converted = sorted([lower * CARBONATE_OR_PHOSPHATE_DT_PER_D18O, upper * CARBONATE_OR_PHOSPHATE_DT_PER_D18O])
                lower, upper = converted[0], converted[1]
                implementation = "Bounded δ18O perturbation converted by locked method-literature local slope −5 °C/‰ before record-wide restandardization"
            else:
                implementation = "Bounded uniform observation-level perturbation in retained measurement units"
            maps.extend(DrawMap(i, "uniform_direct", lower=lower, upper=upper, implementation=implementation, source_layer="Layer 2 bounded conservative") for i in rows)
            active = True

        elif dist == "record_wide_correlated_calibration_realization":
            # The locked table explicitly prohibits a constant offset because it
            # vanishes after record-wide z-standardization. It supplies a residual
            # scale but no raw calibration variable, coefficient covariance, or
            # recalculation equation. A shape-changing realization would therefore
            # require an invented model and is not executed in this strict pilot.
            implementation = "Fixed: table requires recalculation/calibration-parameter perturbation, but raw calibration variables and coefficient covariance are not executable parameters"

        elif dist == "fixed":
            implementation = "Fixed: no defensible quantitative measurement parameter"
        else:
            raise RuntimeError(f"Unsupported Tier 2 distribution {dist} for {rid}")

        audit.append({
            "record_id": rid,
            "n_measurements": len(rows),
            "primary_structure": r.measurement_primary_structure,
            "table_distribution": dist,
            "actually_active": active,
            "actually_perturbed_measurements": len(rows) if active else 0,
            "random_analytical": bool(r.random_analytical),
            "calibration_reconstruction": bool(r.calibration_reconstruction),
            "systematic_upper_bound": bool(r.systematic_upper_bound),
            "execution_implementation": implementation,
            "table_execution_status": status,
        })
    return maps, pd.DataFrame(audit)


def assign_bins(ages: np.ndarray, spec: dict) -> np.ndarray:
    width, lower, count = spec["width"], spec["lower"], spec["count"]
    idx = np.full(len(ages), -1, dtype=int)
    for j, age in enumerate(ages):
        for i in range(count):
            lo = lower + i * width
            hi = lower + (i + 1) * width
            if lo <= age < hi:
                idx[j] = i
                break
    return idx


def iqr_linear(x: np.ndarray) -> float:
    return float(np.quantile(x, 0.75, method="linear") - np.quantile(x, 0.25, method="linear"))


def build_masks(raw: pd.DataFrame, iqr_table: pd.DataFrame, record_code: np.ndarray):
    ages = raw.age_Ma.to_numpy(float)
    base = raw.value.to_numpy(float)
    z = np.empty(len(raw), dtype=float)
    for code in np.unique(record_code):
        ix = np.flatnonzero(record_code == code)
        z[ix] = (base[ix] - base[ix].mean()) / base[ix].std(ddof=1)

    record_lookup = {rid: i for i, rid in enumerate(pd.unique(raw.record_id))}
    masks = {}
    for scheme, spec in BIN_SPECS.items():
        assigned = assign_bins(ages, spec)
        centers = spec["lower"] + (np.arange(spec["count"]) + 0.5) * spec["width"]
        source = iqr_table[
            iqr_table.scheme.astype(str).eq(scheme)
            & iqr_table.alignment.eq("standard")
            & iqr_table.IQR.gt(0)
        ].copy()
        groups = []
        counts = np.zeros(spec["count"], dtype=int)
        sums = np.zeros(spec["count"], dtype=float)
        for _, row in source.iterrows():
            b = int(np.argmin(abs(centers - float(row.bin_center_Ma))))
            ix = np.flatnonzero((record_code == record_lookup[row.record_id]) & (assigned == b))
            if len(ix) != int(row.n_measurements):
                raise RuntimeError(f"Observed-mask count mismatch: {scheme}/{row.record_id}/{row.bin_center_Ma}")
            q = iqr_linear(z[ix])
            if abs(q - float(row.IQR)) > 5e-10:
                raise RuntimeError(f"Observed IQR mismatch: {scheme}/{row.record_id}/{row.bin_center_Ma}: {q} vs {row.IQR}")
            groups.append((b, ix, row.record_id))
            counts[b] += 1
            sums[b] += q
        if np.any(counts == 0):
            raise RuntimeError(f"Observed mask leaves empty bins for {scheme} kyr")
        masks[scheme] = {"centers": centers, "groups": groups, "bin_counts": counts, "observed_ev": sums / counts, "assigned": assigned}
    return masks


def validate_s1_50(masks, l1_module) -> float:
    from openpyxl import load_workbook
    wb = load_workbook(S1_PATH, read_only=True, data_only=True)
    ws = wb["total IQR"]
    source = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) >= 16 and l1_module.numeric(row[14]) and l1_module.numeric(row[15]):
            source[round(float(row[15]) / 1000.0, 8)] = float(row[14])
    diffs = [abs(source[round(float(a), 8)] - float(v)) for a, v in zip(masks["50"]["centers"], masks["50"]["observed_ev"])]
    return max(diffs)


def coverage_table(raw: pd.DataFrame, masks: dict, active_rows: set[int]) -> pd.DataFrame:
    rows = []
    for scheme, info in masks.items():
        per_bin = {i: set() for i in range(len(info["centers"]))}
        for b, ix, _rid in info["groups"]:
            per_bin[b].update(int(x) for x in ix)
        for b, age in enumerate(info["centers"]):
            contributing = per_bin[b]
            active = contributing.intersection(active_rows)
            rows.append({
                "bin_width_kyr": int(scheme),
                "midpoint_Ma": float(age),
                "contributing_measurements": len(contributing),
                "perturbed_contributing_measurements": len(active),
                "perturbation_coverage_proportion": len(active) / len(contributing) if contributing else math.nan,
            })
    return pd.DataFrame(rows).sort_values(["bin_width_kyr", "midpoint_Ma"], ascending=[True, False])


def simulate(raw: pd.DataFrame, mappings: list[DrawMap], masks: dict):
    rng = np.random.default_rng(SEED)
    record_ids = pd.unique(raw.record_id)
    record_lookup = {rid: i for i, rid in enumerate(record_ids)}
    record_code = raw.record_id.map(record_lookup).to_numpy(int)
    record_indices = [np.flatnonzero(record_code == code) for code in range(len(record_ids))]
    base = raw.value.to_numpy(float)

    normal = [m for m in mappings if m.distribution == "normal_direct"]
    uniform = [m for m in mappings if m.distribution == "uniform_direct"]
    isotope = [m for m in mappings if m.distribution == "odp1049_isotope_normal"]
    normal_ix = np.array([m.row_index for m in normal], dtype=int)
    normal_sigma = np.array([m.sigma for m in normal], dtype=float)
    uniform_ix = np.array([m.row_index for m in uniform], dtype=int)
    uniform_lo = np.array([m.lower for m in uniform], dtype=float)
    uniform_hi = np.array([m.upper for m in uniform], dtype=float)
    isotope_ix = np.array([m.row_index for m in isotope], dtype=int)
    isotope_sigma = np.array([m.sigma for m in isotope], dtype=float)
    isotope_base = np.array([m.isotope_value for m in isotope], dtype=float)
    isotope_species = np.array([m.species for m in isotope])

    outputs = {k: np.empty((ITERATIONS, len(v["centers"])), dtype=float) for k, v in masks.items()}
    contrasts = np.empty((ITERATIONS, 3), dtype=float)
    anomaly = {
        "nonfinite_value_iterations": 0,
        "nonfinite_z_iterations": 0,
        "nonfinite_EV_scheme_iterations": 0,
        "zero_or_invalid_record_SD_events": 0,
        "missing_EV_trajectory_iterations": 0,
        "temperature_below_absolute_zero_events": 0,
    }
    max_abs_z = 0.0
    min_ev, max_ev = math.inf, -math.inf
    category_min = {c: math.inf for c in raw.category.unique()}
    category_max = {c: -math.inf for c in raw.category.unique()}

    for it in range(ITERATIONS):
        values = base.copy()
        if len(normal_ix):
            values[normal_ix] += rng.normal(0.0, normal_sigma)
        if len(uniform_ix):
            values[uniform_ix] += rng.uniform(uniform_lo, uniform_hi)
        if len(isotope_ix):
            iso = isotope_base + rng.normal(0.0, isotope_sigma)
            s = isotope_species == "S"
            values[isotope_ix[s]] = 0.03 * iso[s] ** 2 - 4.43 * iso[s] + 10.2855
            values[isotope_ix[~s]] = 0.03 * iso[~s] ** 2 - 4.427 * iso[~s] + 10.064075

        if not np.all(np.isfinite(values)):
            anomaly["nonfinite_value_iterations"] += 1
        temp_values = values[raw.category.to_numpy() == "temperature"]
        anomaly["temperature_below_absolute_zero_events"] += int(np.sum(temp_values < -273.15))
        for cat in category_min:
            v = values[raw.category.to_numpy() == cat]
            category_min[cat] = min(category_min[cat], float(np.min(v)))
            category_max[cat] = max(category_max[cat], float(np.max(v)))

        z = np.empty(len(raw), dtype=float)
        for ix in record_indices:
            sd = values[ix].std(ddof=1)
            if not math.isfinite(sd) or sd <= 0:
                anomaly["zero_or_invalid_record_SD_events"] += 1
                z[ix] = np.nan
            else:
                z[ix] = (values[ix] - values[ix].mean()) / sd
        if not np.all(np.isfinite(z)):
            anomaly["nonfinite_z_iterations"] += 1
        max_abs_z = max(max_abs_z, float(np.nanmax(np.abs(z))))

        for scheme, info in masks.items():
            sums = np.zeros(len(info["centers"]), dtype=float)
            for b, ix, _rid in info["groups"]:
                sums[b] += iqr_linear(z[ix])
            ev = sums / info["bin_counts"]
            outputs[scheme][it] = ev
            if not np.all(np.isfinite(ev)):
                anomaly["nonfinite_EV_scheme_iterations"] += 1
            min_ev = min(min_ev, float(np.nanmin(ev)))
            max_ev = max(max_ev, float(np.nanmax(ev)))
        if any(not np.all(np.isfinite(outputs[k][it])) for k in outputs):
            anomaly["missing_EV_trajectory_iterations"] += 1

        ages = masks["50"]["centers"]
        ev50 = outputs["50"][it]
        phase_mean = {}
        for name, (lo, hi) in PHASES.items():
            sel = (ages >= lo - 1e-10) & (ages <= hi + 1e-10)
            phase_mean[name] = float(ev50[sel].mean())
        contrasts[it] = [
            phase_mean["Phase 1"] - phase_mean["pre-escalation"],
            phase_mean["Phase 2"] - phase_mean["pre-escalation"],
            phase_mean["Recovery"] - phase_mean["Phase 2"],
        ]

    anomaly.update({
        "maximum_absolute_simulated_z_score": max_abs_z,
        "minimum_simulated_global_EV_50kyr": min_ev,
        "maximum_simulated_global_EV_50kyr": max_ev,
        "simulated_measurement_min_by_category": category_min,
        "simulated_measurement_max_by_category": category_max,
    })
    return outputs, contrasts, anomaly


def summarize(raw, params, layer1_maps, layer1_notes, tier2_audit, all_maps, masks, outputs, contrasts, anomaly, coverage, sources, s1_diff, protected_before):
    rows = []
    for scheme, arr in outputs.items():
        q = np.quantile(arr, [0.025, 0.5, 0.975], axis=0)
        info = masks[scheme]
        for age, obs, lo, med, hi in zip(info["centers"], info["observed_ev"], q[0], q[1], q[2]):
            rows.append({
                "bin_width_kyr": int(scheme), "midpoint_Ma": float(age), "observed_EV": float(obs),
                "simulation_median": float(med), "simulation_2.5th": float(lo), "simulation_97.5th": float(hi),
                "interval_width": float(hi - lo), "observed_minus_simulation_median": float(obs - med),
            })
    summary = pd.DataFrame(rows).sort_values(["bin_width_kyr", "midpoint_Ma"], ascending=[True, False])
    summary50 = summary[summary.bin_width_kyr.eq(50)].copy()
    summary50.to_csv(RESULT_DIR / "Layer2_Global_EV_50kyr_summary.csv", index=False)

    names = ["Phase 1 - pre-escalation", "Phase 2 - pre-escalation", "Recovery - Phase 2"]
    cq = np.quantile(contrasts, [0.025, 0.5, 0.975], axis=0)
    contrast = pd.DataFrame({"contrast": names, "simulation_median": cq[1], "simulation_2.5th": cq[0], "simulation_97.5th": cq[2], "interval_width": cq[2] - cq[0]})
    contrast.to_csv(RESULT_DIR / "Layer2_phase_contrasts_50kyr.csv", index=False)
    coverage.to_csv(RESULT_DIR / "Layer2_perturbation_coverage_by_bin.csv", index=False)

    layer1_counts = pd.Series([raw.loc[m.row_index, "record_id"] for m in layer1_maps]).value_counts().to_dict()
    l1_params = params[params.measurement_analysis_layer.eq("documented measurement propagation")]
    l1_audit = []
    for _, r in l1_params.iterrows():
        mapped = int(layer1_counts.get(r.record_id, 0))
        l1_audit.append({
            "record_id": r.record_id, "n_measurements": int(r.n_measurements), "primary_structure": r.measurement_primary_structure,
            "table_distribution": r.measurement_distribution, "actually_active": mapped > 0,
            "actually_perturbed_measurements": mapped, "random_analytical": bool(r.random_analytical),
            "calibration_reconstruction": bool(r.calibration_reconstruction), "systematic_upper_bound": bool(r.systematic_upper_bound),
            "execution_implementation": layer1_notes.get(r.record_id, ""), "table_execution_status": r.measurement_execution_status,
        })
    execution_audit = pd.concat([pd.DataFrame(l1_audit), tier2_audit], ignore_index=True)
    if len(execution_audit) != 70 or execution_audit.record_id.nunique() != 70:
        raise RuntimeError("Execution audit does not map exactly 70 unique records")
    execution_audit.to_csv(RESULT_DIR / "Layer2_record_execution_audit.csv", index=False)

    active_rows = {m.row_index for m in all_maps}
    actual_records = sorted(raw.loc[list(active_rows), "record_id"].unique())
    primary_summary = []
    for label, selector in [
        ("Documented Layer 1A source-mapped", execution_audit.primary_structure.str.startswith("Tier 1")),
        ("Tier 2 conservative Gaussian observation-level", execution_audit.primary_structure.str.contains("Conservative Gaussian")),
        ("Tier 2 bounded uniform/truncated observation-level", execution_audit.primary_structure.str.contains("Bounded uniform")),
        ("Tier 2 record-wide correlated calibration/reconstruction (not executable)", execution_audit.primary_structure.str.contains("Record-wide correlated")),
        ("Fixed: no defensible quantitative basis", execution_audit.primary_structure.str.contains("No defensible")),
    ]:
        part = execution_audit[selector]
        primary_summary.append({
            "summary_type": "mutually exclusive primary execution structure", "structure": label,
            "records_in_parameter_table": len(part), "measurements_in_parameter_table": int(part.n_measurements.sum()),
            "actually_active_records": int(part.actually_active.sum()), "actually_perturbed_measurements": int(part.actually_perturbed_measurements.sum()),
        })
    for flag, label in [("random_analytical", "Random analytical flag"), ("calibration_reconstruction", "Calibration/reconstruction flag"), ("systematic_upper_bound", "Systematic/upper-bound flag")]:
        part = execution_audit[execution_audit[flag] & execution_audit.actually_active]
        primary_summary.append({
            "summary_type": "non-mutually-exclusive active uncertainty flag", "structure": label,
            "records_in_parameter_table": int(execution_audit[execution_audit[flag]].shape[0]),
            "measurements_in_parameter_table": int(execution_audit.loc[execution_audit[flag], "n_measurements"].sum()),
            "actually_active_records": len(part), "actually_perturbed_measurements": int(part.actually_perturbed_measurements.sum()),
        })
    structure_summary = pd.DataFrame(primary_summary)
    structure_summary.to_csv(RESULT_DIR / "Layer2_structure_participation_summary.csv", index=False)

    absdiff = summary50.observed_minus_simulation_median.abs()
    widths = summary50.interval_width
    maxdiff_row = summary50.iloc[int(absdiff.argmax())]
    maxwidth_row = summary50.iloc[int(widths.argmax())]
    protected_after = {str(p): sha256(p) for p in PROTECTED_PATHS if p.exists()}
    unchanged = protected_before == protected_after

    metadata = {
        "analysis": "Layer 2 conservative method-based measurement sensitivity pilot — standard 50-kyr Global EV only",
        "iterations": ITERATIONS,
        "random_seed": SEED,
        "chronology_perturbed": False,
        "retained_measurements": int(len(raw)),
        "retained_operational_records": int(raw.record_id.nunique()),
        "actually_active_records": len(actual_records),
        "actually_perturbed_measurements": len(active_rows),
        "perturbed_measurement_percent": 100 * len(active_rows) / len(raw),
        "documented_active_records": int(((execution_audit.primary_structure.str.startswith("Tier 1")) & execution_audit.actually_active).sum()),
        "documented_perturbed_measurements": int(execution_audit.loc[execution_audit.primary_structure.str.startswith("Tier 1"), "actually_perturbed_measurements"].sum()),
        "tier2_gaussian_active_records": int((execution_audit.primary_structure.str.contains("Conservative Gaussian") & execution_audit.actually_active).sum()),
        "tier2_gaussian_perturbed_measurements": int(execution_audit.loc[execution_audit.primary_structure.str.contains("Conservative Gaussian"), "actually_perturbed_measurements"].sum()),
        "tier2_bounded_active_records": int((execution_audit.primary_structure.str.contains("Bounded uniform") & execution_audit.actually_active).sum()),
        "tier2_bounded_perturbed_measurements": int(execution_audit.loc[execution_audit.primary_structure.str.contains("Bounded uniform"), "actually_perturbed_measurements"].sum()),
        "tier2_correlated_active_records": int((execution_audit.primary_structure.str.contains("Record-wide correlated") & execution_audit.actually_active).sum()),
        "tier2_correlated_perturbed_measurements": int(execution_audit.loc[execution_audit.primary_structure.str.contains("Record-wide correlated"), "actually_perturbed_measurements"].sum()),
        "s1_50kyr_max_absolute_validation_difference": s1_diff,
        "observed_EV_vs_simulation_median_max_absolute_difference": float(absdiff.max()),
        "max_difference_midpoint_Ma": float(maxdiff_row.midpoint_Ma),
        "50kyr_interval_width_median": float(widths.median()),
        "50kyr_interval_width_maximum": float(widths.max()),
        "50kyr_maximum_width_midpoint_Ma": float(maxwidth_row.midpoint_Ma),
        "protected_dataset_and_document_hashes_unchanged": unchanged,
        "anomaly_diagnostics": anomaly,
        "implementation_notes": [
            "Layer 1A documented observation mappings were reused without adding the three unresolved Os records.",
            "Tier 2 Gaussian and bounded observation-level records were perturbed according to the locked table. For temperature targets expressed in underlying δ18O, the method-literature equivalence encoded during feasibility work was implemented as a local −5 °C/‰ conversion before full record-wide restandardization.",
            "The seven record-wide correlated calibration/reconstruction records remained fixed because the locked executable table explicitly prohibits constant additive shifts and does not provide raw calibration variables, equations, or coefficient covariance needed for a shape-changing recalculation.",
            "The seven no-quantitative-basis records remained fixed. Ages, bins, and the observed non-zero record-bin mask were fixed in all realizations.",
        ],
    }
    (RESULT_DIR / "pilot_metadata.json").write_text(json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8")

    manifest = []
    for p in [RAW_PATH, IQR_PATH, S1_PATH, PARAM_PATH, LAYER1_SCRIPT, SCRIPT]:
        manifest.append({"file": p.name, "sha256": sha256(p), "source": str(p)})
    for name, p in sources.items():
        manifest.append({"file": name, "sha256": sha256(p), "source": str(p)})
    pd.DataFrame(manifest).drop_duplicates("sha256").to_csv(RESULT_DIR / "input_source_hash_manifest.csv", index=False)
    pd.DataFrame([{"path": k, "sha256_before": v, "sha256_after": protected_after.get(k), "unchanged": v == protected_after.get(k)} for k, v in protected_before.items()]).to_csv(RESULT_DIR / "protected_files_hash_check.csv", index=False)

    np.savez_compressed(
        RESULT_DIR / "Layer2_Global_EV_realizations_50kyr.npz",
        EV_50kyr=outputs["50"],
        age_50kyr=masks["50"]["centers"],
        phase_contrasts_50kyr=contrasts,
    )
    return metadata, summary50, contrast, execution_audit, structure_summary


def main():
    protected_before = {str(p): sha256(p) for p in PROTECTED_PATHS if p.exists()}
    raw = pd.read_csv(RAW_PATH).reset_index(drop=True)
    raw["record_id"] = raw.record_id.map(lambda x: unicodedata.normalize("NFC", str(x)).strip())
    if len(raw) != 10881 or raw.record_id.nunique() != 70:
        raise RuntimeError(f"Unexpected retained baseline dimensions: {len(raw)} measurements, {raw.record_id.nunique()} records")
    params = parameter_table()
    layer1_maps, layer1_notes, sources = build_layer1_maps(raw)
    tier2, tier2_audit = tier2_maps(raw, params)
    all_maps = layer1_maps + tier2
    row_ids = [m.row_index for m in all_maps]
    if len(row_ids) != len(set(row_ids)):
        raise RuntimeError("An observation received more than one measurement uncertainty mapping")

    record_ids = pd.unique(raw.record_id)
    record_lookup = {rid: i for i, rid in enumerate(record_ids)}
    record_code = raw.record_id.map(record_lookup).to_numpy(int)
    iqr_table = pd.read_csv(IQR_PATH)
    iqr_table["record_id"] = iqr_table.record_id.map(lambda x: unicodedata.normalize("NFC", str(x)).strip())
    masks = build_masks(raw, iqr_table, record_code)
    l1_module = load_layer1_module()
    s1_diff = validate_s1_50(masks, l1_module)
    if s1_diff > 1e-12:
        raise RuntimeError(f"Observed 50-kyr EV does not reproduce Dataset S1 v6: {s1_diff}")

    active_rows = set(row_ids)
    coverage = coverage_table(raw, masks, active_rows)
    outputs, contrasts, anomaly = simulate(raw, all_maps, masks)
    metadata, summary50, contrast, audit, structures = summarize(
        raw, params, layer1_maps, layer1_notes, tier2_audit, all_maps, masks, outputs,
        contrasts, anomaly, coverage, sources, s1_diff, protected_before,
    )
    print(json.dumps(metadata, indent=2, ensure_ascii=False))
    print("\nPHASE CONTRASTS")
    print(contrast.to_string(index=False))
    print("\nPRIMARY/FLAG PARTICIPATION")
    print(structures.to_string(index=False))


if __name__ == "__main__":
    main()
