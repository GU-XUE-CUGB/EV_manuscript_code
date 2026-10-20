#!/usr/bin/env python3
"""Primary record-specific chronological-sensitivity pilot for 50-kyr Global EV.

Locked scope
------------
* 1,000 realizations and the standard 50-kyr reconstruction only.
* Measurement values and record-wide z-scores remain fixed.
* Three independent shared-age-model groups receive one whole-record shift
  each per realization; all operational records within a group share that
  shift exactly.
* No common absolute-timescale systematic is introduced.
* Ages are reassigned to bins in every realization. Record-bin eligibility is
  recomputed dynamically using n >= 2 and IQR > 0; the observed eligibility
  mask is never reused.
* Group-only diagnostics reuse the same group-specific shift draws as the
  primary scenario and therefore require no additional random sampling.
"""

from __future__ import annotations

import hashlib
import json
import math
import unicodedata
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook


MASTER_SEED = 20260824
ITERATIONS = 1000
TOTAL_RETAINED_MEASUREMENTS = 10881
TOTAL_OPERATIONAL_RECORDS = 70

SCRIPT = Path(__file__).resolve()
WORK_ROOT = SCRIPT.parent.parent
DATA_DIR = WORK_ROOT / "01_data_revised"
FORMAL_DIR = DATA_DIR / "formal_datasets"
INVENTORY_DIR = WORK_ROOT / "03_results" / "EV_uncertainty_inventory"
GROUPING_DIR = INVENTORY_DIR / "Chronology_shared_age_model_grouping_audit"
RESULT_DIR = INVENTORY_DIR / "Chronology_primary_sensitivity_pilot_50kyr_1000"
RESULT_DIR.mkdir(parents=True, exist_ok=True)

RAW_PATH = DATA_DIR / "retained_measurements_with_zscores_v4.csv"
IQR_PATH = DATA_DIR / "record_level_IQR_all_bin_widths_revised_v4_n2.csv"
S1_PATH = FORMAL_DIR / "Dataset_S1_revised_v6.xlsx"
GROUPING_PATH = GROUPING_DIR / "Chronology_shared_age_model_grouping_audit.xlsx"

BIN_SPEC = {"width": 0.05, "lower": 64.775, "count": 49}
PHASES = {
    "pre-escalation": (66.95, 67.20),
    "Phase 1": (66.45, 66.90),
    "Phase 2": (66.05, 66.40),
    "Recovery": (64.80, 65.95),
}
CONTRAST_NAMES = [
    "Phase 1 - pre-escalation",
    "Phase 2 - pre-escalation",
    "Recovery - Phase 2",
]
GROUP_ORDER = ["CHR-P1-ODP1209", "CHR-P2-U1403", "CHR-P3-SEYMOUR"]
EXPECTED_GROUPS = {
    "CHR-P1-ODP1209": {"records": 16, "measurements": 2848, "lo_kyr": -87.0, "hi_kyr": 87.0},
    "CHR-P2-U1403": {"records": 3, "measurements": 1249, "lo_kyr": -37.0, "hi_kyr": 37.0},
    "CHR-P3-SEYMOUR": {"records": 5, "measurements": 133, "lo_kyr": -40.0, "hi_kyr": 40.0},
}

PROTECTED_PATHS = [
    GROUPING_PATH,
    S1_PATH,
    FORMAL_DIR / "Dataset_S2_revised.xlsx",
    FORMAL_DIR / "Dataset_S3_revised_v3.xlsx",
    FORMAL_DIR / "Dataset_S4_revised_v3.xlsx",
    FORMAL_DIR / "Dataset_S5_revised_v4.xlsx",
    FORMAL_DIR / "Dataset_S6_revised_v4.xlsx",
    WORK_ROOT / "05_manuscript_revised" / "Manuscript_working_stage3_sampling_density.docx",
    WORK_ROOT / "06_supplement_revised" / "Supplementary_materials_working_stage3_sampling_density.docx",
]


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def clean_id(value) -> str:
    return unicodedata.normalize("NFC", str(value)).strip()


def numeric(value) -> bool:
    return isinstance(value, (int, float)) and math.isfinite(float(value))


def load_grouping_audit() -> tuple[pd.DataFrame, pd.DataFrame]:
    wb = load_workbook(GROUPING_PATH, read_only=True, data_only=True)

    ws = wb["Primary record membership"]
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    headers = [clean_id(x) for x in rows[0]]
    membership = pd.DataFrame(rows[1:], columns=headers).dropna(how="all")
    membership["operational_record"] = membership.operational_record.map(clean_id)
    membership["chronology_group_id"] = membership.chronology_group_id.map(clean_id)
    membership["n_measurements"] = membership.n_measurements.astype(int)

    ws = wb["Primary groups"]
    rows = list(ws.iter_rows(min_row=4, values_only=True))
    headers = [clean_id(x) for x in rows[0]]
    groups = pd.DataFrame(rows[1:], columns=headers).dropna(how="all")
    groups["chronology_group_id"] = groups.chronology_group_id.map(clean_id)
    groups["operational_record_count"] = groups.operational_record_count.astype(int)
    groups["measurement_count"] = groups.measurement_count.astype(int)
    return membership, groups


def validate_grouping(raw: pd.DataFrame, membership: pd.DataFrame, groups: pd.DataFrame) -> dict[str, list[str]]:
    if len(membership) != 24 or membership.operational_record.nunique() != 24:
        raise RuntimeError("Primary chronology membership must contain 24 unique operational records")
    if set(membership.chronology_group_id) != set(GROUP_ORDER):
        raise RuntimeError("Primary chronology group IDs differ from the locked three-group scenario")
    if set(groups.chronology_group_id) != set(GROUP_ORDER):
        raise RuntimeError("Primary groups sheet differs from the locked group IDs")

    raw_counts = raw.groupby("record_id").size()
    group_members: dict[str, list[str]] = {}
    for gid in GROUP_ORDER:
        members = membership.loc[membership.chronology_group_id.eq(gid), "operational_record"].tolist()
        missing = sorted(set(members) - set(raw_counts.index))
        if missing:
            raise RuntimeError(f"{gid} contains records missing from retained raw data: {missing}")
        n_records = len(members)
        n_measurements = int(raw_counts.loc[members].sum())
        expected = EXPECTED_GROUPS[gid]
        if n_records != expected["records"] or n_measurements != expected["measurements"]:
            raise RuntimeError(
                f"{gid} expected {expected['records']} records/{expected['measurements']} measurements; "
                f"found {n_records}/{n_measurements}"
            )
        audit_n = int(membership.loc[membership.chronology_group_id.eq(gid), "n_measurements"].sum())
        if audit_n != n_measurements:
            raise RuntimeError(f"{gid} raw/audit measurement count mismatch: {n_measurements} vs {audit_n}")
        group_members[gid] = members

    if sum(len(v) for v in group_members.values()) != 24:
        raise RuntimeError("Primary membership total is not 24 records")
    if sum(EXPECTED_GROUPS[g]["measurements"] for g in GROUP_ORDER) != 4230:
        raise RuntimeError("Locked primary measurement total is not 4,230")
    return group_members


def assign_bins(ages: np.ndarray) -> np.ndarray:
    """Literal half-open [lower, upper) assignment used by the EV reconstruction."""
    width, lower, count = BIN_SPEC["width"], BIN_SPEC["lower"], BIN_SPEC["count"]
    idx = np.full(len(ages), -1, dtype=np.int16)
    for b in range(count):
        lo = lower + b * width
        hi = lower + (b + 1) * width
        mask = (idx < 0) & (ages >= lo) & (ages < hi)
        idx[mask] = b
    return idx


def iqr_linear(values: np.ndarray) -> float:
    q = np.quantile(values, [0.25, 0.75], method="linear")
    return float(q[1] - q[0])


def prepare_record_arrays(raw: pd.DataFrame) -> dict[str, tuple[np.ndarray, np.ndarray]]:
    result = {}
    for rid, part in raw.groupby("record_id", sort=False):
        ages = part.age_Ma.to_numpy(float)
        z = part.z_score_recomputed.to_numpy(float)
        if not np.all(np.isfinite(ages)) or not np.all(np.isfinite(z)):
            raise RuntimeError(f"Non-finite retained age or z-score in {rid}")
        if len(z) < 2 or abs(float(z.mean())) > 1e-8 or abs(float(z.std(ddof=1)) - 1.0) > 1e-8:
            raise RuntimeError(f"Record-wide z-score validation failed for {rid}")
        result[rid] = (ages, z)
    return result


def record_contributions(
    record_arrays: dict[str, tuple[np.ndarray, np.ndarray]],
    record_ids: list[str],
    shift_ma: float,
) -> tuple[np.ndarray, np.ndarray, dict]:
    """Return the sum/count of eligible record IQRs in every 50-kyr bin."""
    n_bins = BIN_SPEC["count"]
    sums = np.zeros(n_bins, dtype=float)
    counts = np.zeros(n_bins, dtype=np.int16)
    diagnostics = {
        "record_bins_with_n_ge_2": 0,
        "eligible_nonzero_record_bins": 0,
        "zero_iqr_record_bins_excluded": 0,
        "invalid_iqr_record_bins": 0,
        "measurements_assigned_to_50kyr_domain": 0,
    }
    for rid in record_ids:
        ages, z = record_arrays[rid]
        assigned = assign_bins(ages + shift_ma)
        diagnostics["measurements_assigned_to_50kyr_domain"] += int(np.sum(assigned >= 0))
        for b in np.unique(assigned[assigned >= 0]):
            values = z[assigned == b]
            if len(values) < 2:
                continue
            diagnostics["record_bins_with_n_ge_2"] += 1
            value = iqr_linear(values)
            if not math.isfinite(value) or value < 0:
                diagnostics["invalid_iqr_record_bins"] += 1
            elif value == 0:
                diagnostics["zero_iqr_record_bins_excluded"] += 1
            else:
                sums[int(b)] += value
                counts[int(b)] += 1
                diagnostics["eligible_nonzero_record_bins"] += 1
    return sums, counts, diagnostics


def read_s1_observed_50kyr() -> pd.DataFrame:
    wb = load_workbook(S1_PATH, read_only=True, data_only=True)
    ws = wb["total IQR"]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if len(row) >= 16 and numeric(row[14]) and numeric(row[15]):
            rows.append({"midpoint_Ma": float(row[15]) / 1000.0, "Dataset_S1_EV": float(row[14])})
    out = pd.DataFrame(rows).drop_duplicates("midpoint_Ma").sort_values("midpoint_Ma").reset_index(drop=True)
    if len(out) != 49:
        raise RuntimeError(f"Expected 49 standard 50-kyr Global EV values in Dataset S1, found {len(out)}")
    return out


def validate_observed_baseline(
    raw: pd.DataFrame,
    record_arrays: dict[str, tuple[np.ndarray, np.ndarray]],
) -> tuple[np.ndarray, np.ndarray, np.ndarray, float, int]:
    centers = BIN_SPEC["lower"] + (np.arange(BIN_SPEC["count"]) + 0.5) * BIN_SPEC["width"]
    sums, counts, diagnostics = record_contributions(record_arrays, list(record_arrays), 0.0)
    if np.any(counts == 0):
        raise RuntimeError("Bottom-up observed reconstruction contains an empty 50-kyr bin")
    observed = sums / counts

    s1 = read_s1_observed_50kyr()
    s1_values = s1.set_index(s1.midpoint_Ma.round(8)).Dataset_S1_EV
    differences = []
    for age, value in zip(centers, observed):
        key = round(float(age), 8)
        if key not in s1_values.index:
            raise RuntimeError(f"Dataset S1 is missing the {age:.2f}-Ma Global EV midpoint")
        differences.append(abs(float(value) - float(s1_values.loc[key])))
    max_s1_diff = float(max(differences))
    if max_s1_diff > 1e-11:
        raise RuntimeError(f"Bottom-up observed EV differs from Dataset S1; max |difference|={max_s1_diff}")

    # Independent validation against the retained record-IQR table.
    iqr = pd.read_csv(IQR_PATH)
    mask = (
        iqr.scheme.astype(str).eq("50")
        & iqr.alignment.eq("standard")
        & iqr.IQR.gt(0)
    )
    submitted = iqr.loc[mask].groupby("bin_center_Ma").IQR.agg(["sum", "count"]).sort_index()
    table_ev = submitted["sum"] / submitted["count"]
    record_iqr_diff = max(
        abs(float(value) - float(table_ev.loc[round(float(age), 8)]))
        for age, value in zip(centers, observed)
    )
    if record_iqr_diff > 1e-11:
        raise RuntimeError(f"Bottom-up observed EV differs from retained record-IQR table: {record_iqr_diff}")
    return centers, observed, counts, max_s1_diff, diagnostics["zero_iqr_record_bins_excluded"]


def draw_group_shifts() -> tuple[dict[str, np.ndarray], dict]:
    seed_sequence = np.random.SeedSequence(MASTER_SEED)
    children = seed_sequence.spawn(len(GROUP_ORDER))
    shifts = {}
    seed_info = {}
    for gid, child in zip(GROUP_ORDER, children):
        rng = np.random.default_rng(child)
        spec = EXPECTED_GROUPS[gid]
        shifts[gid] = rng.uniform(spec["lo_kyr"], spec["hi_kyr"], size=ITERATIONS)
        seed_info[gid] = {
            "master_entropy": int(child.entropy),
            "spawn_key": list(child.spawn_key),
            "bit_generator": rng.bit_generator.__class__.__name__,
            "draw_min_kyr": float(shifts[gid].min()),
            "draw_max_kyr": float(shifts[gid].max()),
        }
    return shifts, seed_info


def phase_contrasts(ev: np.ndarray, ages: np.ndarray) -> np.ndarray:
    phase_means = {}
    for name, (lo, hi) in PHASES.items():
        select = (ages >= lo - 1e-10) & (ages <= hi + 1e-10)
        phase_means[name] = np.nanmean(ev[:, select], axis=1)
    return np.column_stack([
        phase_means["Phase 1"] - phase_means["pre-escalation"],
        phase_means["Phase 2"] - phase_means["pre-escalation"],
        phase_means["Recovery"] - phase_means["Phase 2"],
    ])


def summarize_ev(scenario: str, ev: np.ndarray, ages: np.ndarray, observed: np.ndarray) -> pd.DataFrame:
    q = np.nanquantile(ev, [0.025, 0.5, 0.975], axis=0)
    rows = []
    for age, obs, lo, med, hi in zip(ages, observed, q[0], q[1], q[2]):
        rows.append({
            "scenario": scenario,
            "midpoint_Ma": float(age),
            "observed_EV": float(obs),
            "simulation_median": float(med),
            "simulation_2.5th": float(lo),
            "simulation_97.5th": float(hi),
            "envelope_width": float(hi - lo),
            "observed_minus_simulation_median": float(obs - med),
            "valid_realizations": int(np.isfinite(ev[:, len(rows)]).sum()),
        })
    return pd.DataFrame(rows)


def summarize_contrasts(scenario: str, contrasts: np.ndarray) -> pd.DataFrame:
    q = np.nanquantile(contrasts, [0.025, 0.5, 0.975], axis=0)
    return pd.DataFrame([
        {
            "scenario": scenario,
            "contrast": name,
            "simulation_median": float(q[1, i]),
            "simulation_2.5th": float(q[0, i]),
            "simulation_97.5th": float(q[2, i]),
            "interval_width": float(q[2, i] - q[0, i]),
            "valid_realizations": int(np.isfinite(contrasts[:, i]).sum()),
        }
        for i, name in enumerate(CONTRAST_NAMES)
    ])


def scenario_overview(scenario: str, summary: pd.DataFrame) -> dict:
    widths = summary.envelope_width.to_numpy(float)
    diff = np.abs(summary.observed_EV - summary.simulation_median).to_numpy(float)
    return {
        "scenario": scenario,
        "median_envelope_width": float(np.nanmedian(widths)),
        "maximum_envelope_width": float(np.nanmax(widths)),
        "maximum_width_midpoint_Ma": float(summary.iloc[int(np.nanargmax(widths))].midpoint_Ma),
        "maximum_observed_median_absolute_difference": float(np.nanmax(diff)),
        "maximum_difference_midpoint_Ma": float(summary.iloc[int(np.nanargmax(diff))].midpoint_Ma),
    }


def main() -> None:
    protected_before = {str(p): sha256(p) for p in PROTECTED_PATHS if p.exists()}
    raw = pd.read_csv(RAW_PATH).reset_index(drop=True)
    raw["record_id"] = raw.record_id.map(clean_id)
    if len(raw) != TOTAL_RETAINED_MEASUREMENTS or raw.record_id.nunique() != TOTAL_OPERATIONAL_RECORDS:
        raise RuntimeError("Final retained baseline must contain 10,881 measurements and 70 records")
    if raw.duplicated(["record_id", "age_Ma", "value", "sample_id"]).any():
        raise RuntimeError("Unexpected exact duplicate retained observations")

    membership, groups = load_grouping_audit()
    group_members = validate_grouping(raw, membership, groups)
    record_arrays = prepare_record_arrays(raw)
    ages, observed, observed_counts, s1_diff, observed_zero_bins = validate_observed_baseline(raw, record_arrays)

    active_records = set(membership.operational_record)
    fixed_records = [rid for rid in record_arrays if rid not in active_records]
    fixed_sums, fixed_counts, fixed_diag = record_contributions(record_arrays, fixed_records, 0.0)
    baseline_group = {}
    for gid in GROUP_ORDER:
        baseline_group[gid] = record_contributions(record_arrays, group_members[gid], 0.0)
    combined_baseline_sums = fixed_sums + sum(baseline_group[g][0] for g in GROUP_ORDER)
    combined_baseline_counts = fixed_counts + sum(baseline_group[g][1] for g in GROUP_ORDER)
    if not np.array_equal(combined_baseline_counts, observed_counts):
        raise RuntimeError("Fixed + group baseline record-bin counts do not reproduce the observed baseline")
    if np.max(np.abs(combined_baseline_sums / combined_baseline_counts - observed)) > 1e-12:
        raise RuntimeError("Fixed + group baseline sums do not reproduce observed Global EV")

    shifts, seed_info = draw_group_shifts()
    group_iteration_sums = {}
    group_iteration_counts = {}
    group_diag_totals = {}
    for gid in GROUP_ORDER:
        sums = np.empty((ITERATIONS, BIN_SPEC["count"]), dtype=float)
        counts = np.empty((ITERATIONS, BIN_SPEC["count"]), dtype=np.int16)
        totals = Counter()
        for iteration, shift_kyr in enumerate(shifts[gid]):
            s, c, d = record_contributions(record_arrays, group_members[gid], shift_kyr / 1000.0)
            sums[iteration] = s
            counts[iteration] = c
            totals.update(d)
        group_iteration_sums[gid] = sums
        group_iteration_counts[gid] = counts
        group_diag_totals[gid] = dict(totals)

    scenarios: dict[str, tuple[np.ndarray, np.ndarray]] = {}
    full_sums = np.broadcast_to(fixed_sums, (ITERATIONS, BIN_SPEC["count"])).copy()
    full_counts = np.broadcast_to(fixed_counts, (ITERATIONS, BIN_SPEC["count"])).copy()
    for gid in GROUP_ORDER:
        full_sums += group_iteration_sums[gid]
        full_counts += group_iteration_counts[gid]
    scenarios["Primary: all three groups"] = (full_sums, full_counts)

    label_by_group = {
        "CHR-P1-ODP1209": "Only CHR-P1-ODP1209 shifted",
        "CHR-P2-U1403": "Only CHR-P2-U1403 shifted",
        "CHR-P3-SEYMOUR": "Only CHR-P3-SEYMOUR shifted",
    }
    for active_gid in GROUP_ORDER:
        sums = np.broadcast_to(fixed_sums, (ITERATIONS, BIN_SPEC["count"])).copy()
        counts = np.broadcast_to(fixed_counts, (ITERATIONS, BIN_SPEC["count"])).copy()
        for gid in GROUP_ORDER:
            if gid == active_gid:
                sums += group_iteration_sums[gid]
                counts += group_iteration_counts[gid]
            else:
                sums += baseline_group[gid][0]
                counts += baseline_group[gid][1]
        scenarios[label_by_group[active_gid]] = (sums, counts)

    ev_arrays = {}
    contrast_arrays = {}
    summaries = []
    contrast_summaries = []
    overview_rows = []
    for scenario, (sums, counts) in scenarios.items():
        ev = np.divide(sums, counts, out=np.full_like(sums, np.nan), where=counts > 0)
        contrasts = phase_contrasts(ev, ages)
        ev_arrays[scenario] = ev
        contrast_arrays[scenario] = contrasts
        summary = summarize_ev(scenario, ev, ages, observed)
        summaries.append(summary)
        contrast_summaries.append(summarize_contrasts(scenario, contrasts))
        overview_rows.append(scenario_overview(scenario, summary))

    all_ev_summary = pd.concat(summaries, ignore_index=True)
    all_contrasts = pd.concat(contrast_summaries, ignore_index=True)
    overview = pd.DataFrame(overview_rows)
    primary_ev = ev_arrays["Primary: all three groups"]
    primary_counts = scenarios["Primary: all three groups"][1]
    primary_summary = all_ev_summary[all_ev_summary.scenario.eq("Primary: all three groups")].copy()
    primary_contrasts = all_contrasts[all_contrasts.scenario.eq("Primary: all three groups")].copy()

    # Peak midpoint in every primary realization; NaN only if an entire trajectory is invalid.
    peak_ages = np.full(ITERATIONS, np.nan)
    for i in range(ITERATIONS):
        if np.any(np.isfinite(primary_ev[i])):
            peak_ages[i] = ages[int(np.nanargmax(primary_ev[i]))]
    peak_counts = pd.Series(peak_ages).value_counts(dropna=False).sort_index()
    peak_frequency = pd.DataFrame({
        "peak_midpoint_Ma": peak_counts.index,
        "frequency": peak_counts.values.astype(int),
        "proportion": peak_counts.values / ITERATIONS,
    })

    eligibility = pd.DataFrame({
        "midpoint_Ma": ages,
        "observed_eligible_record_bins": observed_counts,
        "simulation_record_bins_min": np.min(primary_counts, axis=0),
        "simulation_record_bins_median": np.median(primary_counts, axis=0),
        "simulation_record_bins_max": np.max(primary_counts, axis=0),
        "simulation_record_bins_2.5th": np.quantile(primary_counts, 0.025, axis=0),
        "simulation_record_bins_97.5th": np.quantile(primary_counts, 0.975, axis=0),
        "iterations_with_empty_bin": np.sum(primary_counts == 0, axis=0),
    })
    total_record_bins = primary_counts.sum(axis=1)

    shift_rows = []
    for i in range(ITERATIONS):
        shift_rows.append({"iteration": i + 1, **{f"{gid}_shift_kyr": float(shifts[gid][i]) for gid in GROUP_ORDER}})
    shift_table = pd.DataFrame(shift_rows)

    group_parameter_rows = []
    for gid in GROUP_ORDER:
        g = groups[groups.chronology_group_id.eq(gid)].iloc[0]
        spec = EXPECTED_GROUPS[gid]
        group_parameter_rows.append({
            "chronology_group_id": gid,
            "site": g.site,
            "operational_records": spec["records"],
            "measurements": spec["measurements"],
            "lower_bound_kyr": spec["lo_kyr"],
            "upper_bound_kyr": spec["hi_kyr"],
            "distribution": "bounded uniform",
            "chronology_source_model": g.chronology_source_model,
            "bound_basis": g.bound_basis,
            "perturbation_rule": "One shared whole-record shift per group per realization",
        })
    group_parameters = pd.DataFrame(group_parameter_rows)

    empty_bin_events = int(np.sum(primary_counts == 0))
    invalid_trajectory_iterations = int(np.sum(~np.all(np.isfinite(primary_ev), axis=1)))
    diff = np.abs(primary_summary.observed_EV - primary_summary.simulation_median).to_numpy(float)
    widths = primary_summary.envelope_width.to_numpy(float)
    anomaly = {
        "iterations_with_at_least_one_empty_bin": int(np.sum(np.any(primary_counts == 0, axis=1))),
        "total_empty_bin_events": empty_bin_events,
        "invalid_or_missing_EV_trajectory_iterations": invalid_trajectory_iterations,
        "nonfinite_primary_EV_values": int(np.sum(~np.isfinite(primary_ev))),
        "observed_total_eligible_record_bins": int(observed_counts.sum()),
        "simulated_total_eligible_record_bins_min": int(total_record_bins.min()),
        "simulated_total_eligible_record_bins_median": float(np.median(total_record_bins)),
        "simulated_total_eligible_record_bins_max": int(total_record_bins.max()),
        "maximum_absolute_per_bin_record_count_change": int(np.max(np.abs(primary_counts - observed_counts[None, :]))),
        "observed_zero_IQR_record_bins_excluded": int(observed_zero_bins),
        "active_group_diagnostics_across_iterations": group_diag_totals,
    }

    active_measurements = int(membership.n_measurements.sum())
    metadata = {
        "analysis": "Primary chronological sensitivity pilot — standard 50-kyr Global EV",
        "iterations": ITERATIONS,
        "master_random_seed": MASTER_SEED,
        "group_random_streams": seed_info,
        "measurement_values_perturbed": False,
        "measurement_uncertainty_included": False,
        "common_absolute_timescale_systematic_included": False,
        "ODP1210_included": False,
        "observed_record_bin_mask_reused": False,
        "bin_assignment": "49 standard 50-kyr half-open bins; 64.775 <= age < 67.225 Ma",
        "eligibility": "n_measurements >= 2 and linearly interpolated IQR > 0, recomputed every realization",
        "global_EV": "record-weighted arithmetic mean of dynamically eligible record-bin IQRs",
        "active_chronology_groups": 3,
        "active_operational_records": len(active_records),
        "active_measurements": active_measurements,
        "active_measurement_coverage_fraction": active_measurements / len(raw),
        "active_measurement_coverage_percent": 100 * active_measurements / len(raw),
        "retained_operational_records": int(raw.record_id.nunique()),
        "retained_measurements": int(len(raw)),
        "Dataset_S1_observed_EV_max_absolute_validation_difference": s1_diff,
        "observed_EV_vs_simulation_median_max_absolute_difference": float(np.max(diff)),
        "maximum_difference_midpoint_Ma": float(primary_summary.iloc[int(np.argmax(diff))].midpoint_Ma),
        "envelope_median_width": float(np.median(widths)),
        "envelope_maximum_width": float(np.max(widths)),
        "maximum_width_midpoint_Ma": float(primary_summary.iloc[int(np.argmax(widths))].midpoint_Ma),
        "observed_peak_midpoint_Ma": float(ages[int(np.argmax(observed))]),
        "anomaly_diagnostics": anomaly,
        "group_only_diagnostic": "Uses the same group-specific 1,000 shift draws as the primary scenario; other chronology groups remain at observed ages.",
    }

    # Save machine-readable results before workbook authoring.
    primary_summary.to_csv(RESULT_DIR / "Chronology_primary_EV_50kyr_summary.csv", index=False)
    primary_contrasts.to_csv(RESULT_DIR / "Chronology_primary_phase_contrasts.csv", index=False)
    peak_frequency.to_csv(RESULT_DIR / "Chronology_primary_peak_age_frequency.csv", index=False)
    eligibility.to_csv(RESULT_DIR / "Chronology_primary_eligibility_diagnostics.csv", index=False)
    all_ev_summary.to_csv(RESULT_DIR / "Chronology_group_contribution_EV_summaries.csv", index=False)
    all_contrasts.to_csv(RESULT_DIR / "Chronology_group_contribution_phase_contrasts.csv", index=False)
    overview.to_csv(RESULT_DIR / "Chronology_group_contribution_overview.csv", index=False)
    shift_table.to_csv(RESULT_DIR / "Chronology_group_shift_draws.csv", index=False)
    group_parameters.to_csv(RESULT_DIR / "Chronology_primary_group_parameters.csv", index=False)
    membership.to_csv(RESULT_DIR / "Chronology_primary_record_membership.csv", index=False)

    np.savez_compressed(
        RESULT_DIR / "Chronology_primary_Global_EV_realizations_50kyr_1000.npz",
        age_50kyr=ages,
        observed_EV_50kyr=observed,
        primary_EV_50kyr=primary_ev,
        primary_eligible_record_counts=primary_counts,
        primary_phase_contrasts=contrast_arrays["Primary: all three groups"],
        ODP1209_only_EV_50kyr=ev_arrays["Only CHR-P1-ODP1209 shifted"],
        U1403_only_EV_50kyr=ev_arrays["Only CHR-P2-U1403 shifted"],
        Seymour_only_EV_50kyr=ev_arrays["Only CHR-P3-SEYMOUR shifted"],
        ODP1209_only_phase_contrasts=contrast_arrays["Only CHR-P1-ODP1209 shifted"],
        U1403_only_phase_contrasts=contrast_arrays["Only CHR-P2-U1403 shifted"],
        Seymour_only_phase_contrasts=contrast_arrays["Only CHR-P3-SEYMOUR shifted"],
        ODP1209_shift_kyr=shifts["CHR-P1-ODP1209"],
        U1403_shift_kyr=shifts["CHR-P2-U1403"],
        Seymour_shift_kyr=shifts["CHR-P3-SEYMOUR"],
        peak_midpoint_Ma=peak_ages,
    )

    protected_after = {str(p): sha256(p) for p in PROTECTED_PATHS if p.exists()}
    protected_check = pd.DataFrame([
        {
            "path": path,
            "sha256_before": digest,
            "sha256_after": protected_after.get(path),
            "unchanged": digest == protected_after.get(path),
        }
        for path, digest in protected_before.items()
    ])
    protected_check.to_csv(RESULT_DIR / "protected_files_hash_check.csv", index=False)
    if not protected_check.unchanged.all():
        raise RuntimeError("At least one protected Dataset/document/source file changed during the pilot")

    manifest_paths = [RAW_PATH, IQR_PATH, S1_PATH, GROUPING_PATH, SCRIPT]
    pd.DataFrame([
        {"file": p.name, "path": str(p), "sha256": sha256(p)} for p in manifest_paths
    ]).to_csv(RESULT_DIR / "input_source_hash_manifest.csv", index=False)
    (RESULT_DIR / "pilot_metadata.json").write_text(json.dumps(metadata, indent=2, ensure_ascii=False), encoding="utf-8")

    print(json.dumps(metadata, indent=2, ensure_ascii=False))
    print("\nPRIMARY PHASE CONTRASTS")
    print(primary_contrasts.to_string(index=False))
    print("\nGROUP-CONTRIBUTION OVERVIEW")
    print(overview.to_string(index=False))
    print("\nALL SCENARIO PHASE CONTRASTS")
    print(all_contrasts.to_string(index=False))


if __name__ == "__main__":
    main()
