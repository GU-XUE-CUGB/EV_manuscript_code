#!/usr/bin/env python3
"""Code 7: reproduce final latitude summaries and Fig. 5/S14/S15."""

import csv
import hashlib
import json
import subprocess
import sys
from collections import Counter, defaultdict
from pathlib import Path

import numpy as np
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent
DATA = ROOT / "inputs"
OUT = ROOT / "output"
OUT.mkdir(exist_ok=True)
SEED = 20260929
ITERATIONS = 10000
AGES = [round(67.2 - 0.05 * i, 2) for i in range(49)]
PHASES = {
    "Pre": [round(67.2 - 0.05 * i, 2) for i in range(6)],
    "Phase 1": [round(66.9 - 0.05 * i, 2) for i in range(10)],
    "Phase 2": [round(66.4 - 0.05 * i, 2) for i in range(8)],
}
SOURCES = {
    "iqr": DATA / "record_level_IQR_all_bin_widths_revised_v4_n2.csv",
    "latitude": DATA / "latitude_record_mapping_v5.csv",
    "canonical": DATA / "canonical_site_mapping_v5.csv",
    "s1": DATA / "Dataset_S1_revised_v6.xlsx",
}

if "--plot-only" in sys.argv:
    subprocess.run(["Rscript", str(ROOT / "plot_latitude_figures_final.R")], check=True)
    raise SystemExit(0)


def rows(path):
    with path.open(newline="", encoding="utf-8-sig") as handle:
        yield from csv.DictReader(handle)


lat = {r["record_id"]: r["latitude_group"] for r in rows(SOURCES["latitude"])}
canonical = defaultdict(set)
for row in rows(SOURCES["canonical"]):
    canonical[row["record_id"]].add(row["canonical_site_v5"])

record_bins = []
for row in rows(SOURCES["iqr"]):
    if row["bin_width_kyr"] != "50" or row["alignment"] != "standard":
        continue
    age = round(float(row["bin_center_Ma"]), 2)
    iqr = float(row["IQR"])
    rid = row["record_id"]
    if age not in AGES or iqr <= 0 or rid == "pCO2::all data":
        continue
    if rid not in lat or len(canonical[rid]) != 1:
        raise ValueError(f"Missing or ambiguous canonical site: {rid}")
    record_bins.append(
        {"age": age, "group": lat[rid], "site": next(iter(canonical[rid])),
         "record_id": rid, "category": row["category"], "iqr": iqr}
    )

# Confirm that the long-table input is exactly the positive latitude IQR pool in final S1.
workbook = load_workbook(SOURCES["s1"], read_only=True, data_only=True)
for group, sheet_name in [
    ("low", "IQR (low latitudes_50 bin)"),
    ("mid_high", "IQR (mid-high latitudes_50 bin)"),
]:
    source = defaultdict(list)
    for row in record_bins:
        if row["group"] == group:
            source[row["age"]].append(row["iqr"])
    target = defaultdict(list)
    for row in workbook[sheet_name].iter_rows(min_row=2, max_col=3, values_only=True):
        if isinstance(row[1], (float, int)) and isinstance(row[2], (float, int)):
            target[round(row[2] / 1000, 2)].append(float(row[1]))
    if set(source) != set(target):
        raise ValueError(f"Bin mismatch against S1: {sheet_name}")
    for age in source:
        a, b = sorted(source[age]), sorted(target[age])
        if len(a) != len(b) or any(abs(x - y) > 1e-8 for x, y in zip(a, b)):
            raise ValueError(f"IQR mismatch against S1: {sheet_name}, {age}")
workbook.close()
if len(record_bins) != 898:
    raise ValueError(f"Expected 898 latitude record-bins, got {len(record_bins)}")

# One site contributes one within-site mean, irrespective of its proxy-record count.
site_iqrs = defaultdict(list)
site_records = defaultdict(set)
site_categories = defaultdict(set)
site_groups = defaultdict(set)
for row in record_bins:
    key = (row["age"], row["group"], row["site"])
    site_iqrs[key].append(row["iqr"])
    site_records[key].add(row["record_id"])
    site_categories[key].add(row["category"])
    site_groups[row["site"]].add(row["group"])
if any(len(groups) != 1 for groups in site_groups.values()):
    raise ValueError("A canonical site belongs to both latitude groups")

contributions = []
by_bin_group = defaultdict(dict)
for (age, group, site), values in sorted(site_iqrs.items(), reverse=True):
    site_ev = float(np.mean(values))
    by_bin_group[(age, group)][site] = site_ev
    contributions.append({
        "age_Ma": age, "latitude_group": group, "canonical_site": site,
        "record_count": len(values), "site_mean_IQR": site_ev,
        "record_ids": "; ".join(sorted(site_records[(age, group, site)])),
        "categories": "; ".join(sorted(site_categories[(age, group, site)])),
    })


def sample_without_replacement(values, k, rng):
    values = np.asarray(values, dtype=float)
    if len(values) == k:
        return np.full(ITERATIONS, values.mean())
    # Ordering by independent uniform keys gives every size-k subset equal probability.
    keys = rng.random((ITERATIONS, len(values)))
    indices = np.argpartition(keys, kth=k - 1, axis=1)[:, :k]
    return values[indices].mean(axis=1)


def quantiles(values):
    q = np.quantile(values, [0.025, 0.5, 0.975], method="linear")
    return {"median": float(q[1]), "p2_5": float(q[0]), "p97_5": float(q[2])}


def simulate(remove_site, seed_sequence):
    rng = np.random.default_rng(seed_sequence)
    draws = {}
    counts = {}
    for age in AGES:
        low = {s: x for s, x in by_bin_group[(age, "low")].items() if s != remove_site}
        high = {s: x for s, x in by_bin_group[(age, "mid_high")].items() if s != remove_site}
        n_low, n_high = len(low), len(high)
        eligible = n_low >= 2 and n_high >= 2
        k = min(n_low, n_high) if eligible else 0
        counts[age] = (n_low, n_high, k, eligible)
        if eligible:
            delta = sample_without_replacement(list(high.values()), k, rng)
            delta -= sample_without_replacement(list(low.values()), k, rng)
            if not np.isfinite(delta).all():
                raise ValueError(f"Non-finite ΔEV at {age}, removed site={remove_site}")
            draws[age] = delta
    return draws, counts


sites = sorted(site_groups)
seed_sequences = np.random.SeedSequence(SEED).spawn(len(sites) + 1)
baseline, baseline_counts = simulate(None, seed_sequences[0])
phase_summary = []
phase_draws = {}
for phase, ages in PHASES.items():
    keep = [age for age in ages if age in baseline]
    arr = np.mean([baseline[age] for age in keep], axis=0) if keep else None
    if arr is not None:
        phase_draws[phase] = arr
    phase_summary.append({
        "phase": phase, "eligible_bins": len(keep), "total_bins": len(ages),
        "eligible_midpoints_Ma": "; ".join(f"{age:.2f}" for age in keep),
        **(quantiles(arr) if arr is not None else {"median": None, "p2_5": None, "p97_5": None}),
    })

bin_summary = []
for age in AGES:
    n_low, n_high, k, eligible = baseline_counts[age]
    row = {
        "midpoint_Ma": age, "low_sites": n_low, "mid_high_sites": n_high,
        "k": k, "eligible": eligible,
        "low_site_names": "; ".join(sorted(by_bin_group[(age, "low")])),
        "mid_high_site_names": "; ".join(sorted(by_bin_group[(age, "mid_high")])),
    }
    if eligible:
        row.update(quantiles(baseline[age]))
    else:
        row.update({"median": None, "p2_5": None, "p97_5": None})
    bin_summary.append(row)

loo_summary = []
for index, site in enumerate(sites):
    loo, _ = simulate(site, seed_sequences[index + 1])
    for phase, ages in PHASES.items():
        keep = [age for age in ages if age in loo]
        if keep:
            arr = np.mean([loo[age] for age in keep], axis=0)
            comparable = np.mean([baseline[age] for age in keep], axis=0)
            estimate = quantiles(arr)
            baseline_matched = float(np.median(comparable))
        else:
            estimate = {"median": None, "p2_5": None, "p97_5": None}
            baseline_matched = None
        loo_summary.append({
            "removed_site": site, "latitude_group": next(iter(site_groups[site])),
            "phase": phase, "eligible_bins_after_removal": len(keep),
            "eligible_midpoints_Ma": "; ".join(f"{age:.2f}" for age in keep),
            **estimate,
            "baseline_same_bins_median": baseline_matched,
            "change_vs_same_bins": estimate["median"] - baseline_matched if keep else None,
        })

descriptive = []
for age in AGES:
    row = {"Age_Ma": age}
    for group in ("low", "mid_high"):
        vals = [r["iqr"] for r in record_bins if r["age"] == age and r["group"] == group]
        row[group + "_EV"] = float(np.mean(vals)) if vals else None
        row[group + "_records"] = len(vals)
        row[group + "_sites"] = len(by_bin_group[(age, group)])
    descriptive.append(row)

def write_csv(name, records):
    with (OUT / name).open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)

write_csv("descriptive_EV.csv", descriptive)
write_csv("site_balanced_bin_summary.csv", bin_summary)
write_csv("site_balanced_phase_summary.csv", phase_summary)
write_csv("leave_one_site_out_summary.csv", loo_summary)
results = {
    "metadata": {
        "seed": SEED, "iterations": ITERATIONS,
        "rng": "NumPy Generator(PCG64), SeedSequence.spawn(1 baseline + one per sorted canonical site)",
        "draw_rule": "Each eligible bin: choose k=min(low sites, mid-high sites) separately without replacement; average site-mean record IQR; ΔEV=mid-high−low",
        "phase_rule": "Arithmetic mean of eligible bin ΔEV in each predefined phase, per iteration; unavailable if no eligible bin",
        "site_count": len(sites), "record_bin_count": len(record_bins),
        "source_sha256": {name: hashlib.sha256(path.read_bytes()).hexdigest() for name, path in SOURCES.items()},
        "canonical_site_note": "Use existing v5 canonical site mapping. Pooled pCO2 is not in the latitude IQR sheets; legacy site aliases within one operational record count as one canonical site.",
    },
    "phase_summary": phase_summary,
    "bin_summary": bin_summary,
    "loo_summary": loo_summary,
    "site_contributions": contributions,
}
(OUT / "site_balanced_rarefaction_results.json").write_text(
    json.dumps(results, indent=2, ensure_ascii=False), encoding="utf-8")
if "--no-plots" not in sys.argv:
    subprocess.run(["Rscript", str(ROOT / "plot_latitude_figures_final.R")], check=True)
print(f"Saved summaries and figures: {OUT}")
